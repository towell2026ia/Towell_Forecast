"""Read-only complete sheet inventory; private context, no database connection."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook  # noqa: E402
from services.assistant_api.historical_corpus import sha256_file  # noqa: E402


def inventory(paths):
    result, seen = [], set()
    for path in paths:
        digest = sha256_file(path)
        if digest in seen:
            result.append({"source_hash": digest, "filename": path.name, "duplicate": True})
            continue
        seen.add(digest)
        book = load_workbook(path, read_only=True, data_only=False)
        sheets = []
        try:
            for sheet in book:
                top, headers, formulas, literals, nonempty = [], [], 0, 0, 0
                sample_rows = []
                for row_no, cells in enumerate(sheet.iter_rows(), 1):
                    populated = {cell.coordinate: cell.value for cell in cells if cell.value is not None}
                    if row_no <= 20:
                        top.append({"row": row_no, "cells": populated})
                    labels = {cell.coordinate: cell.value for cell in cells if isinstance(cell.value, str)
                              and not cell.value.startswith("=")}
                    if any(str(label).strip().casefold() in {"item", "upc", "prime item nbr", "gral", "cadena", "código", "codigo"}
                           for label in labels.values()):
                        headers.append({"row": row_no, "labels": labels})
                    if len(sample_rows) < 3 and any(isinstance(cell.value, (int, float)) and cell.value >= 10000 for cell in cells[:5]):
                        sample_rows.append({"row": row_no, "cells": populated})
                    for cell in cells:
                        if cell.value is not None:
                            nonempty += 1
                            if cell.data_type == "f" or hasattr(cell.value, "text"):
                                formulas += 1
                            elif isinstance(cell.value, (int, float)):
                                literals += 1
                sheets.append({"sheet_name": sheet.title, "range": sheet.calculate_dimension(),
                               "rows": sheet.max_row, "columns": sheet.max_column, "formula_count": formulas,
                               "numeric_literal_count": literals, "nonempty": nonempty, "headers": headers,
                               "top_context": top, "identity_sample_context": sample_rows})
        finally:
            book.close()
        result.append({"source_hash": digest, "filename": path.name, "duplicate": False, "sheets": sheets})
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, action="append", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    path = args.output.resolve()
    if not path.is_relative_to(Path(__file__).resolve().parents[1] / "outputs"):
        parser.error("private output must be in ignored outputs/")
    if path.exists():
        parser.error("preserve existing inventory; use a fresh path")
    data = inventory(args.source)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, default=str)
    print(json.dumps([{"file": row["filename"], "duplicate": row["duplicate"],
                       "sheets": len(row.get("sheets", []))} for row in data], indent=2))


if __name__ == "__main__": main()
