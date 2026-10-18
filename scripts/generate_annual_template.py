"""Generate the empty annual commercial template using Python/openpyxl only.

No data is fetched, imported, forecast or written to a database.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from .annual_template_contract import ROOT, SCHEMA, SHEETS, confirmed_rows, headers
except ImportError:
    from annual_template_contract import ROOT, SCHEMA, SHEETS, confirmed_rows, headers

NAVY, SLATE, AMBER, GRAY = "183B56", "63778B", "FFF5D6", "EFF3F6"


def _validation(ws, letter: str, column: dict) -> None:
    cell = f"{letter}2"
    kind = column["type"]
    kwargs = {"type": "custom", "formula1": f'OR({cell}="",AND(ISTEXT({cell}),NOT(ISFORMULA({cell}))))'}
    if "enum" in column:
        kwargs = {"type": "list", "formula1": '"' + ",".join(column["enum"]) + '"'}
    elif kind == "quantity":
        kwargs["formula1"] = f'OR({cell}="",AND(ISNUMBER({cell}),{cell}>=0,{cell}<10^15,ROUND({cell},3)={cell},NOT(ISFORMULA({cell}))))'
    elif kind == "integer":
        minimum = 2024 if column["name"] == "anio" else 1
        maximum = 9999 if column["name"] == "anio" else 2147483647
        kwargs = {"type": "whole", "operator": "between", "formula1": minimum, "formula2": maximum}
    elif kind == "date":
        kwargs = {"type": "date", "operator": "between", "formula1": "DATE(1900,1,1)", "formula2": "DATE(9999,12,31)"}
    elif kind == "period":
        year = f',LEFT({cell},4)=TEXT(DeclaredYear,"0000")' if column["name"] == "periodo" else ""
        kwargs["formula1"] = (f'OR({cell}="",IFERROR(AND(ISTEXT({cell}),LEN({cell})=7,MID({cell},5,1)="-",'
                              f'TEXT(DATE(VALUE(LEFT({cell},4)),VALUE(RIGHT({cell},2)),1),"yyyy-mm")={cell}{year}),FALSE))')
    rule = DataValidation(**kwargs, allow_blank=True)
    rule.showErrorMessage = True
    rule.errorTitle = "Revisar formato"
    rule.error = "Consulte 00_INSTRUCCIONES. No sustituya vacíos por cero ni utilice fórmulas en captura."
    rule.showInputMessage = True
    rule.promptTitle = "Obligatorio" if column["required"] else "Opcional"
    rule.prompt = column["description"][:250]
    rule.errorStyle = "stop"
    ws.add_data_validation(rule)
    rule.add(f"{letter}2:{letter}{2 if ws.title == '01_CONTROL_CARGA' else 1048576}")


def _table(ws, name: str, end: int) -> None:
    table = Table(displayName=SHEETS[name]["table"], ref=f"A1:{get_column_letter(len(headers(name)))}{max(2, end)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def _instructions(ws) -> None:
    rows = [
        ["Objetivo", "Un archivo por cadena y año", "Obligatorio", "Incluya todos los modelos y productos de una cadena. Años desde 2024. Plantilla local; no realiza cargas productivas.", "CADENA_2024_v1.xlsx"],
        ["Orden", "1. Fuentes y control", "Obligatorio", "Registre las fuentes en 05_FUENTES y una sola fila en 01_CONTROL_CARGA. fuente_principal referencia una fuente existente.", "fuente_id documental"],
        ["Orden", "2. Productos", "Obligatorio", "Una fila por código de origen + variante. No deduzca identidades por descripción, ITEM o UPC.", "Código textual y variante"],
        ["Orden", "3. Histórico", "Obligatorio", "Capture sólo meses con evidencia. Al menos una cantidad real por fila. No agregue doce meses automáticamente.", "2024-01"],
        ["Orden", "4. Forecast Cliente", "Opcional", "Conserve versiones, emisión, objetivo y recepción. No capture cálculos C1/C2 o Champion.", "Emisión 2024-12, objetivo 2025-01"],
        ["Orden", "5. Diccionarios y validación", "Consulta", "Revise equivalencias explícitas y ejecute Python. 07_VALIDACIONES tiene prechequeos; --report-workbook completa una copia de consulta.", "Ver GUIDE_CARGA_ANUAL.md"],
        ["Cantidades", "Vacío y cero", "Regla", "Vacío = dato no proporcionado. 0 = cantidad real confirmada como cero. No invente datos ni modifique cantidades para pasar validaciones.", "0 / celda vacía"],
        ["Identidad", "ITEM y UPC", "Regla", "Ingrese como texto antes de pegar. El validador rechaza números aunque se aplique formato texto después. Conserve ceros iniciales.", "0001234567890 (texto)"],
        ["Estado", "Clasificación comercial", "Regla", "Homologar una escritura no aprueba un producto. Nuevos modelos, colores y presentaciones son aceptados sin listas cerradas.", "PENDING_CLASSIFICATION"],
        ["Fuentes", "Fechas documentales", "Regla", "El corte de una fuente no demuestra disponibilidad histórica. Deje vacías las fechas desconocidas. Los hashes los calcula Python con --source-dir.", "YYYY-MM-DD"],
        ["Colores", "Campos obligatorios", "Leyenda", "Encabezado azul oscuro y primera fila ámbar: obligatorio cuando existe un registro. No cambie los encabezados.", "Celda de captura"],
        ["Colores", "Opcionales y calculados", "Leyenda", "Encabezado gris azulado: opcional. 07_VALIDACIONES es gris y de consulta. No se rellena manualmente.", "Sin columnas ocultas"],
        ["Escala", "Tablas de Excel", "Regla", "Pegue dentro de la tabla o use Tab desde su última celda para ampliarla. Las validaciones cubren las filas de Excel. Python lee también filas fuera de la tabla.", "No añada hojas ni columnas"],
        ["Validación", "Resultado completo", "Regla", "Python produce validation_summary.json, validation_errors.csv y validation_status.json. --report-workbook genera una copia, nunca modifica el original.", "VALIDO / ADVERTENCIA / ERROR"],
        ["Errores", "Problemas frecuentes", "Consulta", "Códigos numéricos, meses con otro año, duplicados, cantidades negativas o fórmulas, fuentes faltantes, BATA en tamaño y Forecast sin emisión.", "Corregir en la fuente con evidencia"],
    ]
    for name, sheet in SHEETS.items():
        if sheet["role"] not in {"input", "optional_input", "rules"}:
            continue
        for col in sheet["columns"]:
            rows.append([name, col["name"], "Obligatorio" if col["required"] else "Opcional", col["description"], col["type"]])
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 33
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 86
    ws.column_dimensions["E"].width = 37
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 56
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _checks(ws) -> None:
    local = "Ejecutar validador local"
    rows = [
        ["Cadena declarada", '=IF(\'01_CONTROL_CARGA\'!A2="","Sin declarar",\'01_CONTROL_CARGA\'!A2)', '=IF(\'01_CONTROL_CARGA\'!A2="","ERROR","ADVERTENCIA")', "Prechequeo. Python valida estructura y relaciones."],
        ["Año", '=IF(\'01_CONTROL_CARGA\'!C2="","Sin declarar",\'01_CONTROL_CARGA\'!C2)', '=IF(\'01_CONTROL_CARGA\'!C2="","ERROR","ADVERTENCIA")', "Año del histórico del archivo."],
        ["Productos registrados", '=COUNTA(T_Productos[codigo_producto_origen])', "ADVERTENCIA", "Preconteo de códigos. El resultado integral lo calcula Python."],
        ["Modelos registrados", local, "ADVERTENCIA", "Conteo único con NFC, mayúsculas y espacios; sin inferir equivalencias."],
        ["Colores registrados", local, "ADVERTENCIA", "Conteo único de etiquetas; sin fusionar colores diferentes."],
        ["Presentaciones registradas", local, "ADVERTENCIA", "Sólo las equivalencias confirmadas se usan para comparar."],
        ["Meses con información", local, "ADVERTENCIA", "Meses con al menos una cantidad proporcionada; 0 cuenta como dato."],
        ["Registros históricos", '=COUNTA(T_Historico[codigo_producto_origen])', "ADVERTENCIA", "Preconteo de códigos. No crea ni completa meses."],
        ["Registros con error", local, "ADVERTENCIA", "Filas distintas con uno o más errores."],
        ["Duplicados", local, "ADVERTENCIA", "Filas adicionales por clave duplicada, no cantidad de mensajes."],
        ["Productos sin identidad suficiente", local, "ADVERTENCIA", "Falta código de origen o descripción literal."],
        ["Productos sin clasificación", local, "ADVERTENCIA", "Pendientes, sin estado o con atributos comerciales incompletos."],
        ["Fuentes no identificadas", local, "ADVERTENCIA", "Referencias vacías o que no existen en 05_FUENTES."],
        ["Estado general", "ADVERTENCIA", "ADVERTENCIA", "Ejecute Python. Un prechequeo de Excel no es validación integral."],
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 33
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 90
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 35
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=GRAY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    for state, color in [("ERROR", "FCE8E6"), ("ADVERTENCIA", AMBER), ("VALIDO", "EEF4EF")]:
        ws.conditional_formatting.add("C2:C15", FormulaRule(formula=[f'C2="{state}"'], fill=PatternFill("solid", fgColor=color)))


def generate(output: Path | str = ROOT / "templates" / SCHEMA["filename"]) -> Path:
    output = Path(output).resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("El archivo de salida debe ser .xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "FORECAST Towell"
    wb.properties.title = "Plantilla de carga anual comercial"
    wb.properties.subject = "Un archivo por cadena y año desde 2024. Sin datos productivos."
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(headers(name))
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C2" if name in {"02_PRODUCTOS", "03_HISTORICO_MENSUAL", "04_FORECAST_CLIENTE"} else "A2"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = SLATE if name in {"00_INSTRUCCIONES", "07_VALIDACIONES"} else NAVY
        ws.print_title_rows = "1:1"
        ws.sheet_view.zoomScale = 90
        ws.row_dimensions[1].height = 44
    wb.defined_names.add(DefinedName("DeclaredYear", attr_text="'01_CONTROL_CARGA'!$C$2"))
    for name, sheet in SHEETS.items():
        ws = wb[name]
        if name == "00_INSTRUCCIONES":
            _instructions(ws)
        elif name == "07_VALIDACIONES":
            _checks(ws)
        else:
            if name == "06_DICCIONARIOS":
                for row in confirmed_rows():
                    ws.append(row)
            for i, col in enumerate(sheet["columns"], 1):
                letter = get_column_letter(i)
                ws.column_dimensions[letter].width = max(20, min(38, len(col["name"]) + 3))
                if col["name"] in {"descripcion_original", "observaciones", "origen_documental"}:
                    ws.column_dimensions[letter].width = 45
                fmt = '#,##0.###' if col["type"] == "quantity" else '0' if col["type"] == "integer" else 'yyyy-mm-dd' if col["type"] == "date" else '@'
                ws.column_dimensions[letter].number_format = fmt
                for row in range(2, 32):
                    cell = ws.cell(row, i)
                    cell.number_format = fmt
                    cell.font = Font(name="Arial", size=11, color="174E8C")
                    cell.alignment = Alignment(vertical="center", wrap_text=col["name"] in {"descripcion_original", "observaciones"})
                    cell.protection = Protection(locked=False)
                    if row == 2:
                        cell.fill = PatternFill("solid", fgColor=AMBER if col["required"] else "FFFFFF")
                _validation(ws, letter, col)
            end = 10 if name == "06_DICCIONARIOS" else 2
            if name == "06_DICCIONARIOS":
                ws.column_dimensions["A"].width = 32
                for row in ws.iter_rows(min_row=2, max_row=10):
                    ws.row_dimensions[row[0].row].height = 36
                    for cell in row:
                        cell.alignment = Alignment(vertical="center", wrap_text=True)
            _table(ws, name, end)
        if name in {"00_INSTRUCCIONES", "07_VALIDACIONES"}:
            _table(ws, name, ws.max_row)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                if name in {"00_INSTRUCCIONES", "07_VALIDACIONES"}:
                    cell.font = Font(name="Arial", size=11, color="25364A")
        for i, col in enumerate(sheet["columns"], 1):
            cell = ws.cell(1, i)
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY if col["required"] else SLATE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "templates" / SCHEMA["filename"])
    args = parser.parse_args()
    output = generate(args.output)
    print(json.dumps({"file": str(output), "sheets": len(SHEETS), "business_records": 0, "schema_version": SCHEMA["schema_version"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
