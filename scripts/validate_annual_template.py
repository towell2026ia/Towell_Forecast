"""Validate an annual Excel file locally; write reports, never change the input.

No connection, credential, database or forecasting dependency is used.
Exit codes: 0 VALIDO/ADVERTENCIA, 1 ERROR, 2 invalid command/output destination.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import uuid
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Protection

try:
    from .annual_template_contract import AMBIGUOUS, CHECK_FORMULAS, INPUT_SHEETS, SCHEMA, SHEETS, blank, comparison_key, headers
except ImportError:
    from annual_template_contract import AMBIGUOUS, CHECK_FORMULAS, INPUT_SHEETS, SCHEMA, SHEETS, blank, comparison_key, headers

PERIOD = re.compile(r"^[0-9]{4}-(0[1-9]|1[0-2])$")
CHAIN = re.compile(r"^[A-Z0-9][A-Z0-9_-]*$")
ERROR_COLUMNS = ["estado", "codigo", "hoja", "fila", "columna", "mensaje", "valor_original", "valor_homologado"]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _period(value: object) -> bool:
    return isinstance(value, str) and PERIOD.fullmatch(value) is not None and 1 <= int(value[:4]) <= 9999


def _date(value: object) -> bool:
    if isinstance(value, datetime):
        return value.time().isoformat() == "00:00:00"
    if isinstance(value, date):
        return True
    if not isinstance(value, str) or not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value):
        return False
    try:
        date.fromisoformat(value)
        return True
    except ValueError:
        return False


def package_problems(path: Path) -> list[str]:
    """Check the OOXML container without executing formulas or following links."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
            return ["ARCHIVO_DEMASIADO_GRANDE"]
        problems = []
        if any("vbaproject" in n.lower() or "macrosheet" in n.lower() for n in names):
            problems.append("MACROS_NO_PERMITIDAS")
        if any(n.startswith("xl/externalLinks/") or n == "xl/connections.xml" or n.startswith("xl/queryTables/") for n in names):
            problems.append("CONEXIONES_NO_PERMITIDAS")
        for name in names:
            if name.endswith(".rels"):
                root = ET.fromstring(archive.read(name))
                if any(rel.get("TargetMode") == "External" for rel in root):
                    problems.append("VINCULO_EXTERNO_NO_PERMITIDO")
                    break
        return problems


class Audit:
    def __init__(self) -> None:
        self.issues: list[dict] = []
        self.rows: dict[str, list[tuple[int, dict]]] = {name: [] for name in INPUT_SHEETS}
        self.missing: Counter = Counter()
        self.differences: Counter = Counter()
        self.duplicates = 0
        self.source_missing = 0
        self.insufficient_identity: set[int] = set()
        self.unclassified: set[int] = set()
        self.source_hashes: list[dict] = []

    def issue(self, code: str, sheet: str, row: int, column: str, message: str,
              state: str = "ERROR", original: str = "", normalized: str = "") -> None:
        self.issues.append(dict(zip(ERROR_COLUMNS, [state, code, sheet, row, column, message, original, normalized])))

    def required(self, sheet: str, row: int, col: str) -> None:
        self.missing[f"{sheet}.{col}"] += 1
        self.issue("CAMPO_OBLIGATORIO", sheet, row, col, "Campo obligatorio sin dato.")

    def read(self, wb) -> None:
        expected = list(SHEETS)
        if wb.sheetnames != expected:
            self.issue("HOJAS_NO_COINCIDEN", "LIBRO", 0, "", "Deben existir exactamente las ocho hojas en el orden del esquema.")
        for name, sheet in SHEETS.items():
            if name not in wb.sheetnames:
                self.issue("HOJA_FALTANTE", name, 1, "", "Falta una hoja requerida.")
                continue
            ws = wb[name]
            first = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            while first and first[-1] is None:
                first.pop()
            if first != headers(name):
                self.issue("ENCABEZADOS_NO_COINCIDEN", name, 1, "", "Nombres, orden o cantidad de columnas diferentes del esquema.")
                continue
            if name not in INPUT_SHEETS:
                for record in ws.iter_rows(min_row=2):
                    for cell in record:
                        if cell.data_type == "f" and (name != "07_VALIDACIONES" or cell.value not in CHECK_FORMULAS):
                            self.issue("FORMULA_NO_PERMITIDA", name, cell.row, cell.column_letter, "Fórmula fuera del área de consulta local permitida.")
                continue
            for record in ws.iter_rows(min_row=2):
                if all(blank(cell.value) for cell in record):
                    continue
                row = record[0].row
                values = {col["name"]: record[i].value for i, col in enumerate(sheet["columns"])}
                self.rows[name].append((row, values))
                if any(not blank(c.value) for c in record[len(sheet["columns"]):]):
                    self.issue("COLUMNA_EXTRA", name, row, "", "Hay datos fuera de las columnas del contrato.")
                for i, column in enumerate(sheet["columns"]):
                    col, kind, cell = column["name"], column["type"], record[i]
                    value = cell.value
                    if blank(value):
                        if column["required"]:
                            self.required(name, row, col)
                        continue
                    if cell.data_type == "f":
                        self.issue("FORMULA_NO_PERMITIDA", name, row, col, "Las hojas de captura no admiten fórmulas.")
                        continue
                    if kind in {"text", "uuid", "period"} and (not isinstance(value, str) or cell.data_type == "e"):
                        self.issue("IDENTIFICADOR_NUMERICO" if col in {"item", "upc", "codigo_producto_origen", "variante_codigo"} else "TIPO_TEXTO", name, row, col, "Debe almacenarse como texto; aplicar formato texto después no recupera ceros.")
                        continue
                    if col in {"codigo_producto_origen", "variante_codigo", "fuente_id", "fuente_principal"} and value != value.strip():
                        self.issue("IDENTIDAD_CON_ESPACIOS", name, row, col, "La clave tiene espacios exteriores. Revisar sin corregirla automáticamente.")
                    if "enum" in column and value not in column["enum"]:
                        self.issue("ESTADO_NO_ADMITIDO", name, row, col, "Valor fuera de los estados definidos en el esquema.")
                    if kind == "uuid":
                        try:
                            if str(uuid.UUID(value)) != value.lower():
                                raise ValueError()
                        except ValueError:
                            self.issue("UUID_INVALIDO", name, row, col, "El UUID opcional tiene formato inválido.")
                        else:
                            self.issue("UUID_NO_CONCILIADO_LOCALMENTE", name, row, col, "El formato es válido; la existencia y correspondencia requieren conciliación externa.", "ADVERTENCIA")
                    elif kind == "integer":
                        minimum = 2024 if col == "anio" else 1
                        maximum = 9999 if col == "anio" else 2147483647
                        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value != int(value) or not minimum <= value <= maximum:
                            self.issue("ANIO_INVALIDO" if col == "anio" else "VERSION_INVALIDA", name, row, col, f"Se requiere entero entre {minimum} y {maximum}.")
                    elif kind == "date" and not _date(value):
                        self.issue("FECHA_INVALIDA", name, row, col, "Se requiere fecha real YYYY-MM-DD o fecha Excel sin hora.")
                    elif kind == "period" and not _period(value):
                        self.issue("PERIODO_INVALIDO", name, row, col, "Se requiere mes textual YYYY-MM válido.")
                    elif kind == "quantity":
                        self.quantity(value, name, row, col)

    def quantity(self, value, sheet: str, row: int, col: str) -> None:
        if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
            self.issue("CANTIDAD_NO_NUMERICA", sheet, row, col, "Se requiere celda numérica; no se convierten textos ni booleanos.")
            return
        try:
            number = Decimal(str(value))
            if not number.is_finite():
                raise InvalidOperation()
            if number < 0:
                self.issue("CANTIDAD_NEGATIVA", sheet, row, col, "Las cantidades no pueden ser negativas.")
            if abs(number) > Decimal(SCHEMA["max_quantity"]):
                self.issue("CANTIDAD_FUERA_DE_RANGO", sheet, row, col, "La cantidad no cabe en numeric(18,3).")
            if number != number.quantize(Decimal("0.001")):
                self.issue("MAS_DE_TRES_DECIMALES", sheet, row, col, "No redondear automáticamente. Revisar la cantidad fuente.")
            if len(number.normalize().as_tuple().digits) > 15:
                self.issue("PRECISION_EXCEL", sheet, row, col, "Revisar precisión de más de 15 dígitos contra el archivo fuente.", "ADVERTENCIA")
        except (InvalidOperation, ValueError, OverflowError):
            self.issue("CANTIDAD_NO_FINITA", sheet, row, col, "La cantidad debe ser finita y representable.")

    def unique(self) -> None:
        duplicate_rows = set()
        for sheet, keys in SCHEMA["unique_keys"].items():
            for key in keys:
                seen = set()
                for row, record in self.rows[sheet]:
                    values = tuple("" if blank(record[c]) else record[c] for c in key)
                    if blank(values[0]):
                        continue
                    if values in seen:
                        duplicate_rows.add((sheet, row))
                        self.issue("DUPLICADO", sheet, row, "+".join(key), "Clave repetida. No se fusionan ni suman registros.")
                    seen.add(values)
        self.duplicates = len(duplicate_rows)

    def reference(self, sheet: str, row: int, value, field: str, sources: set) -> None:
        if blank(value) or value not in sources:
            self.source_missing += 1
            self.issue("FUENTE_INEXISTENTE", sheet, row, field, "Falta la referencia o no existe en 05_FUENTES.")

    def relationships(self) -> None:
        controls = self.rows["01_CONTROL_CARGA"]
        if len(controls) != 1:
            self.issue("CONTROL_CARGA_UNICO", "01_CONTROL_CARGA", 2, "", "Debe existir exactamente una fila de control: una cadena y un año.")
            if not controls:
                for col in SHEETS["01_CONTROL_CARGA"]["columns"]:
                    if col["required"]:
                        self.required("01_CONTROL_CARGA", 2, col["name"])
        control = controls[0][1] if len(controls) == 1 else {}
        chain = control.get("codigo_cadena")
        if not isinstance(chain, str) or not CHAIN.fullmatch(chain):
            self.issue("CADENA_INVALIDA", "01_CONTROL_CARGA", 2, "codigo_cadena", "Declare un único código de cadena verificado con letras mayúsculas, números, guion o guion bajo.")
        sources = {r["fuente_id"] for _, r in self.rows["05_FUENTES"] if isinstance(r["fuente_id"], str) and not blank(r["fuente_id"])}
        for row, record in controls:
            self.reference("01_CONTROL_CARGA", row, record["fuente_principal"], "fuente_principal", sources)
        products = {(r["codigo_producto_origen"], r["variante_codigo"] or "") for _, r in self.rows["02_PRODUCTOS"]}
        if not self.rows["02_PRODUCTOS"]:
            self.issue("CATALOGO_VACIO", "02_PRODUCTOS", 2, "codigo_producto_origen", "No hay productos registrados.")
        if not self.rows["03_HISTORICO_MENSUAL"]:
            self.issue("HISTORICO_VACIO", "03_HISTORICO_MENSUAL", 2, "", "El archivo no aporta observaciones históricas.", "ADVERTENCIA")
        for sheet in ["02_PRODUCTOS", "03_HISTORICO_MENSUAL", "04_FORECAST_CLIENTE"]:
            for row, record in self.rows[sheet]:
                self.reference(sheet, row, record["fuente_id"], "fuente_id", sources)
                if sheet != "02_PRODUCTOS" and (record["codigo_producto_origen"], record["variante_codigo"] or "") not in products:
                    self.issue("PRODUCTO_INEXISTENTE", sheet, row, "codigo_producto_origen+variante_codigo", "La clave exacta no existe en 02_PRODUCTOS. No se buscan coincidencias por nombre.")
        year = control.get("anio")
        for row, record in self.rows["03_HISTORICO_MENSUAL"]:
            period = record["periodo"]
            if _period(period) and isinstance(year, (int, float)) and not isinstance(year, bool) and int(period[:4]) != year:
                self.issue("PERIODO_FUERA_DEL_ANIO", "03_HISTORICO_MENSUAL", row, "periodo", "El mes debe pertenecer al año declarado en CONTROL_CARGA.")
            if all(blank(record[col]) for col in ("venta_real", "pedido", "enviado")):
                self.issue("CANTIDADES_AUSENTES", "03_HISTORICO_MENSUAL", row, "venta_real+pedido+enviado", "La fila no contiene ninguna cantidad real. No crear registros vacíos.")
        versions = {}
        for row, record in self.rows["04_FORECAST_CLIENTE"]:
            issue, target = record["periodo_emision"], record["periodo_objetivo"]
            if blank(issue):
                self.issue("FORECAST_SIN_EMISION", "04_FORECAST_CLIENTE", row, "periodo_emision", "Forecast Cliente sin periodo de emisión.")
            if _period(issue) and _period(target) and target < issue:
                self.issue("FORECAST_RETROSPECTIVO", "04_FORECAST_CLIENTE", row, "periodo_objetivo", "Objetivo anterior a emisión. Conservar evidencia y revisar; no se convierte en forecast oportuno.", "ADVERTENCIA")
            if blank(record["estado_documento"]):
                self.issue("ESTADO_DOCUMENTO_AUSENTE", "04_FORECAST_CLIENTE", row, "estado_documento", "El estado documental no fue proporcionado. No se asigna uno automáticamente.", "ADVERTENCIA")
            key = (issue, record["version_forecast_cliente"])
            metadata = (str(record["fecha_recepcion"]), record["estado_documento"])
            if key in versions and versions[key] != metadata:
                self.issue("VERSION_CLIENTE_INCONSISTENTE", "04_FORECAST_CLIENTE", row, "fecha_recepcion+estado_documento", "Una misma emisión y versión tiene metadatos incompatibles entre productos.")
            versions[key] = metadata

    def commercial(self) -> None:
        pending = set()
        equivalences = {}
        for row, record in self.rows["06_DICCIONARIOS"]:
            attr, original, normalized = record["atributo"], record["valor_original"], record["valor_homologado"]
            if not isinstance(original, str):
                continue
            key = (attr, comparison_key(original))
            if record["estado_equivalencia"] == "CONFIRMADA":
                if blank(normalized):
                    self.required("06_DICCIONARIOS", row, "valor_homologado")
                elif isinstance(normalized, str):
                    if key in equivalences and equivalences[key] != comparison_key(normalized):
                        self.issue("EQUIVALENCIA_CONTRADICTORIA", "06_DICCIONARIOS", row, "valor_homologado", "El mismo valor tiene equivalencias confirmadas incompatibles.")
                    equivalences[key] = comparison_key(normalized)
                    if [attr, original, normalized] not in SCHEMA["confirmed_equivalences"]:
                        self.issue("REGLA_ADICIONAL_POR_REVISAR", "06_DICCIONARIOS", row, "valor_original", "Regla adicional requiere revisión; el validador no cambia atributos.", "ADVERTENCIA")
            elif record["estado_equivalencia"] == "PENDIENTE":
                pending.add(key)
                self.issue("EQUIVALENCIA_PENDIENTE", "06_DICCIONARIOS", row, "estado_equivalencia", "Equivalencia pendiente; no se aplica.", "ADVERTENCIA")
        confirmed_sizes = {comparison_key(a): b for attr, a, b in SCHEMA["confirmed_equivalences"] if attr == "TAMANO_PRESENTACION"}
        for row, record in self.rows["02_PRODUCTOS"]:
            if any(blank(record[c]) for c in ["codigo_producto_origen", "descripcion_original"]):
                self.insufficient_identity.add(row)
            if record["estado_clasificacion"] != "APPROVED" or any(blank(record[c]) for c in ["categoria", "tipo_producto", "modelo", "color", "tamano_presentacion"]):
                self.unclassified.add(row)
                self.issue("CLASIFICACION_PENDIENTE", "02_PRODUCTOS", row, "estado_clasificacion", "Clasificación provisional o incompleta. No se promueve a APPROVED.", "ADVERTENCIA")
            if record["estado_clasificacion"] == "APPROVED":
                self.issue("APROBACION_NO_CONCILIADA", "02_PRODUCTOS", row, "estado_clasificacion", "Estado declarado; la evidencia y autorización no se verifican offline ni se aplican.", "ADVERTENCIA")
            for field, attr in [("modelo", "MODELO"), ("color", "COLOR"), ("tamano_presentacion", "TAMANO_PRESENTACION"), ("tipo_producto", "TIPO_PRODUCTO")]:
                original = record[field]
                if not isinstance(original, str) or blank(original):
                    if field == "modelo":
                        self.issue("MODELO_SIN_CLASIFICACION", "02_PRODUCTOS", row, field, "Modelo no proporcionado; el producto permanece visible y pendiente.", "ADVERTENCIA")
                    continue
                key = comparison_key(original)
                if key in AMBIGUOUS or (attr, key) in pending:
                    self.issue("COLOR_AMBIGUO" if field == "color" else "TAMANO_AMBIGUO" if field == "tamano_presentacion" else "ATRIBUTO_AMBIGUO", "02_PRODUCTOS", row, field, "Etiqueta ambigua o equivalencia pendiente. No inferir un valor.", "ADVERTENCIA")
                normalized = comparison_key(original) if field == "modelo" else confirmed_sizes.get(key, original) if field == "tamano_presentacion" else "BATA" if field == "tipo_producto" and key in {"BATA", "BATAS"} else original
                if field == "tamano_presentacion" and key in {"BATA", "BATAS"}:
                    self.issue("BATA_ES_TIPO_NO_TAMANO", "02_PRODUCTOS", row, field, "BATA/BATAS corresponde a tipo_producto. No se mueve automáticamente.")
                if original != normalized:
                    self.differences[field] += 1
                    self.issue("HOMOLOGACION_PROPUESTA", "02_PRODUCTOS", row, field, "Diferencia de escritura informativa; el archivo original no se modifica.", "ADVERTENCIA", original, normalized)

    def hash_sources(self, directory: Path | None) -> None:
        if directory is None:
            return
        base = directory.resolve()
        cache = {}
        for row, record in self.rows["05_FUENTES"]:
            filename = record["nombre_archivo"]
            if not isinstance(filename, str) or blank(filename):
                continue
            path = (base / filename).resolve()
            if not path.is_relative_to(base) or not path.is_file():
                self.issue("ARCHIVO_FUENTE_NO_LOCALIZADO", "05_FUENTES", row, "nombre_archivo", "Fuente no localizada dentro del directorio autorizado.", "ADVERTENCIA")
                continue
            if path not in cache:
                cache[path] = sha256_file(path)
            self.source_hashes.append({"fuente_id": record["fuente_id"], "nombre_archivo": filename, "sha256": cache[path]})

    def result(self) -> dict:
        def count_labels(field, normalize=lambda s: s):
            return len({normalize(r[field]) for _, r in self.rows["02_PRODUCTOS"] if isinstance(r[field], str) and not blank(r[field])})
        sizes = {comparison_key(a): b for attr, a, b in SCHEMA["confirmed_equivalences"] if attr == "TAMANO_PRESENTACION"}
        months = {r["periodo"] for _, r in self.rows["03_HISTORICO_MENSUAL"] if _period(r["periodo"]) and any(not blank(r[c]) for c in ["venta_real", "pedido", "enviado"])}
        counts = Counter(i["estado"] for i in self.issues)
        state = "ERROR" if counts["ERROR"] else "ADVERTENCIA" if counts["ADVERTENCIA"] else "VALIDO"
        control = self.rows["01_CONTROL_CARGA"][0][1] if len(self.rows["01_CONTROL_CARGA"]) == 1 else {}
        return {"schema_version": SCHEMA["schema_version"], "estado_general": state,
                "cadena_declarada": control.get("codigo_cadena"), "anio": control.get("anio"),
                "registros": {k: len(v) for k, v in self.rows.items()},
                "productos_registrados": len(self.rows["02_PRODUCTOS"]), "modelos_registrados": count_labels("modelo", comparison_key),
                "colores_registrados": count_labels("color", comparison_key), "presentaciones_registradas": count_labels("tamano_presentacion", lambda v: sizes.get(comparison_key(v), comparison_key(v))),
                "meses_con_informacion": len(months), "registros_historicos": len(self.rows["03_HISTORICO_MENSUAL"]),
                "errores": counts["ERROR"], "advertencias": counts["ADVERTENCIA"],
                "registros_con_error": len({(i["hoja"], i["fila"]) for i in self.issues if i["estado"] == "ERROR"}),
                "duplicados": self.duplicates, "productos_sin_identidad_suficiente": len(self.insufficient_identity),
                "productos_sin_clasificacion": len(self.unclassified), "fuentes_no_identificadas": self.source_missing,
                "campos_faltantes": dict(sorted(self.missing.items())), "diferencias": dict(self.differences)}


def validate(path: Path | str, source_dir: Path | None = None) -> tuple[dict, list[dict], dict]:
    path = Path(path).resolve()
    audit = Audit()
    before = sha256_file(path)
    try:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("FORMATO_NO_XLSX")
        problems = package_problems(path)
        for code in problems:
            audit.issue(code, "LIBRO", 0, "", "El contenedor debe ser .xlsx sin macros ni conexiones externas.")
        if not problems:
            wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                audit.read(wb)
            finally:
                wb.close()
            audit.unique()
            audit.relationships()
            audit.commercial()
            audit.hash_sources(source_dir)
    except (ValueError, TypeError, KeyError, OSError, ET.ParseError, zipfile.BadZipFile, InvalidOperation):
        audit.issue("ARCHIVO_NO_VALIDABLE", "LIBRO", 0, "", "Archivo ilegible, formato no admitido o estructura inválida. Revisar la fuente.")
    after = sha256_file(path)
    if before != after:
        audit.issue("ARCHIVO_CAMBIO_DURANTE_LECTURA", "LIBRO", 0, "", "El archivo cambió durante la validación. Repetir con una copia estable.")
    summary = audit.result()
    status = {"estado": summary["estado_general"], "schema_version": SCHEMA["schema_version"],
              "input_sha256": before, "input_sha256_after": after, "original_sin_modificar": before == after,
              "validated_at_utc": datetime.now(timezone.utc).isoformat(), "supabase_connected": False,
              "productive_import_implemented": False, "source_hashes": audit.source_hashes}
    return summary, audit.issues, status


def _safe_csv(value) -> str:
    value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")) else value


def write_report_copy(source: Path, target: Path, summary: dict, status: dict) -> None:
    if target.resolve() == source.resolve() or target.exists() and target.samefile(source):
        raise ValueError("El informe no puede sobrescribir el original.")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("La copia de consulta debe ser .xlsx.")
    wb = load_workbook(source, data_only=False, keep_links=False)
    try:
        # Display rejected formulas literally in the separate report; never execute them.
        for sheet in wb:
            for record in sheet:
                for cell in record:
                    if cell.data_type == "f":
                        cell.data_type = "s"
        ws = wb["07_VALIDACIONES"]
        labels = ["cadena_declarada", "anio", "productos_registrados", "modelos_registrados", "colores_registrados", "presentaciones_registradas", "meses_con_informacion", "registros_historicos", "registros_con_error", "duplicados", "productos_sin_identidad_suficiente", "productos_sin_clasificacion", "fuentes_no_identificadas", "estado_general"]
        for row, key in enumerate(labels, 2):
            ws.cell(row, 2, summary[key] if summary[key] is not None else "Sin declarar")
            if isinstance(summary[key], str):
                ws.cell(row, 2).data_type = "s"
            ws.cell(row, 3, summary["estado_general"])
            ws.cell(row, 4, f"Resultado Python. Archivo fuente SHA-256 {status['input_sha256']}. No editar esta copia de consulta.")
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    cell.protection = Protection(locked=True)
            sheet.protection.sheet = True
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
    finally:
        wb.close()


def write_reports(source: Path, output_dir: Path, summary: dict, issues: list[dict], status: dict, report_workbook: Path | None = None) -> None:
    targets = [output_dir / name for name in ["validation_summary.json", "validation_errors.csv", "validation_status.json"]]
    if report_workbook:
        targets.append(report_workbook)
    for target in targets:
        if target.resolve() == source.resolve() or target.exists() and target.samefile(source):
            raise ValueError("Los informes no pueden sobrescribir el original.")
    if len({p.resolve() for p in targets}) != len(targets):
        raise ValueError("Las rutas de informes deben ser diferentes.")
    output_dir.mkdir(parents=True, exist_ok=True)
    targets[0].write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with targets[1].open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=ERROR_COLUMNS)
        writer.writeheader()
        writer.writerows({k: _safe_csv(v) for k, v in issue.items()} for issue in issues)
    targets[2].write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if report_workbook:
        # Never repackage a workbook rejected for macros, links or corruption.
        if package_problems(source) or "07_VALIDACIONES" not in load_workbook_sheetnames(source):
            raise ValueError("No se puede generar copia de consulta de un contenedor inválido.")
        write_report_copy(source, report_workbook, summary, status)


def load_workbook_sheetnames(source: Path) -> list[str]:
    wb = load_workbook(source, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path.cwd() / "outputs/annual_template_validation")
    parser.add_argument("--source-dir", type=Path, help="Directorio local de fuentes reales para calcular SHA-256; no usa red.")
    parser.add_argument("--report-workbook", type=Path, help="Copia opcional con 07_VALIDACIONES completa; nunca el original.")
    args = parser.parse_args()
    if not args.file.is_file():
        parser.error("No existe el archivo de entrada.")
    try:
        summary, issues, status = validate(args.file, args.source_dir)
        write_reports(args.file.resolve(), args.output_dir, summary, issues, status, args.report_workbook)
    except (ValueError, OSError, zipfile.BadZipFile):
        parser.error("No fue posible escribir los informes en rutas distintas del original.")
    print(json.dumps(summary, ensure_ascii=False))
    raise SystemExit(1 if summary["estado_general"] == "ERROR" else 0)


if __name__ == "__main__":
    main()
