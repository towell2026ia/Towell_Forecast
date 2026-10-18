"""Isolated synthetic workbooks only. No network, database or engine calls."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from annual_template_contract import SCHEMA, SHEETS, comparison_key, headers  # noqa: E402
from check_annual_template_public import OFFICIAL, audit_template  # noqa: E402
from generate_annual_template import generate  # noqa: E402
from validate_annual_template import sha256_file, validate, write_reports  # noqa: E402


def set_row(wb, sheet, row=2, **values):
    for key, value in values.items():
        wb[sheet].cell(row, headers(sheet).index(key) + 1, value)


def synthetic(wb, year=2024, model="FENDI OXFORD"):
    set_row(wb, "01_CONTROL_CARGA", codigo_cadena="SYNTH_CHAIN", nombre_cadena="Cadena sintética",
            anio=year, version_archivo="v1", responsable="Prueba aislada", fecha_corte_fuente=date(year, 12, 31),
            fuente_principal="SYNTH_SOURCE")
    set_row(wb, "05_FUENTES", fuente_id="SYNTH_SOURCE", nombre_archivo="synthetic.csv", tipo_fuente="CSV",
            origen_documental="Generado para pruebas", responsable_fuente="Prueba aislada")
    set_row(wb, "02_PRODUCTOS", codigo_producto_origen="000001", item="000012", upc="0000000012345",
            descripcion_original="Descripción sintética literal", categoria="SYNTH_CATEGORY", tipo_producto="TOALLA",
            modelo=model, color="AZUL", tamano_presentacion="MB", estado_clasificacion="PENDING_CLASSIFICATION",
            fuente_id="SYNTH_SOURCE")
    set_row(wb, "03_HISTORICO_MENSUAL", codigo_producto_origen="000001", periodo=f"{year}-01",
            venta_real=0, pedido=12.125, fuente_id="SYNTH_SOURCE")


class AnnualTemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.file = generate(self.base / "synthetic.xlsx")

    def edit(self, callback=None, year=2024):
        wb = load_workbook(self.file)
        synthetic(wb, year)
        if callback:
            callback(wb)
        wb.save(self.file)
        wb.close()
        return validate(self.file)

    def codes(self, result):
        return {i["codigo"] for i in result[1]}

    def test_official_is_empty_exact_schema_and_safe(self):
        self.assertEqual(audit_template(ROOT / OFFICIAL), [])
        self.assertEqual(audit_template(self.file), [])
        wb = load_workbook(self.file)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, list(SHEETS))
        for name in SHEETS:
            self.assertEqual([c.value for c in wb[name][1]], headers(name))
            self.assertEqual(len(wb[name].tables), 1)
            self.assertIsNotNone(wb[name].freeze_panes)
            self.assertFalse(any(c.hidden for c in wb[name].column_dimensions.values()))
        self.assertEqual(wb["02_PRODUCTOS"]["C2"].number_format, "@")
        self.assertEqual(wb["02_PRODUCTOS"]["D2"].number_format, "@")
        self.assertGreater(len(wb["02_PRODUCTOS"].data_validations.dataValidation), 0)
        model_rule = [r for r in wb["02_PRODUCTOS"].data_validations.dataValidation if "H2" in str(r.sqref)][0]
        self.assertEqual(model_rule.type, "custom")
        self.assertIn("1048576", str(model_rule.sqref))

    def test_empty_fails_for_capture_but_never_modified(self):
        before = sha256_file(self.file)
        summary, issues, status = validate(self.file)
        self.assertEqual(summary["estado_general"], "ERROR")
        self.assertEqual(summary["productos_registrados"], 0)
        self.assertIn("CADENA_INVALIDA", {i["codigo"] for i in issues})
        self.assertTrue(status["original_sin_modificar"])
        self.assertEqual(before, sha256_file(self.file))

    def test_years_incomplete_months_blank_and_real_zero(self):
        for year in [2024, 2025, 2026]:
            with self.subTest(year=year):
                summary, _, status = self.edit(year=year)
                self.assertEqual(summary["errores"], 0)
                self.assertEqual(summary["meses_con_informacion"], 1)
                self.assertEqual(summary["registros_historicos"], 1)
                self.assertTrue(status["original_sin_modificar"])
                wb = load_workbook(self.file)
                self.assertEqual(wb["03_HISTORICO_MENSUAL"]["D2"].value, 0)
                self.assertIsNone(wb["03_HISTORICO_MENSUAL"]["F2"].value)
                wb.close()

    def test_models_are_open_unicode_no_fuzzy_merging(self):
        self.assertEqual(comparison_key("  Fendi   Oxford "), "FENDI OXFORD")
        self.assertEqual(comparison_key("BAn\u0303O"), "BAÑO")
        self.assertNotEqual(comparison_key("FENDI"), comparison_key("FENDI OXFORD"))
        def add(wb):
            for row, model in enumerate(["Fendi Oxford", " FENDI  OXFORD ", "MODELO NUEVO", "FIORE", "GRANADA"], 3):
                set_row(wb, "02_PRODUCTOS", row, codigo_producto_origen=f"NEW{row}", descripcion_original="Mismo nombre no implica identidad",
                        modelo=model, tipo_producto="BATAS" if row == 6 else "TOALLA", color="NUEVO COLOR",
                        tamano_presentacion=["PULL", "FACIAL", "PRESENTACION NUEVA", "MEDIO BAÑO", "BANO"][row - 3],
                        fuente_id="SYNTH_SOURCE")
        result = self.edit(add)
        self.assertEqual(result[0]["errores"], 0)
        self.assertEqual(result[0]["modelos_registrados"], 4)
        self.assertGreater(result[0]["diferencias"]["modelo"], 0)

    def test_optional_existing_uuid_and_leading_zero_identity(self):
        result = self.edit(lambda wb: set_row(wb, "02_PRODUCTOS", producto_id_supabase="00000000-0000-4000-8000-000000000001"))
        self.assertEqual(result[0]["errores"], 0)
        self.assertIn("UUID_NO_CONCILIADO_LOCALMENTE", self.codes(result))
        wb = load_workbook(self.file)
        self.assertEqual(wb["02_PRODUCTOS"]["A2"].value, "000001")
        self.assertEqual(wb["02_PRODUCTOS"]["D2"].value, "0000000012345")
        wb.close()

    def test_numeric_item_and_upc_rejected_even_if_text_format(self):
        result = self.edit(lambda wb: set_row(wb, "02_PRODUCTOS", item=12, upc=12345))
        issues = [i for i in result[1] if i["codigo"] == "IDENTIFICADOR_NUMERICO"]
        self.assertEqual({i["columna"] for i in issues}, {"item", "upc"})

    def test_invalid_control_and_multiple_chains(self):
        for value in [2023, "2024", 2024.5, 10000]:
            result = self.edit(lambda wb: set_row(wb, "01_CONTROL_CARGA", anio=value))
            self.assertIn("ANIO_INVALIDO", self.codes(result))
        result = self.edit(lambda wb: set_row(wb, "01_CONTROL_CARGA", 3, codigo_cadena="OTHER", anio=2024))
        self.assertIn("CONTROL_CARGA_UNICO", self.codes(result))

    def test_history_period_and_year(self):
        for period, code in [("2024-13", "PERIODO_INVALIDO"), ("2024-1", "PERIODO_INVALIDO"), ("2025-01", "PERIODO_FUERA_DEL_ANIO")]:
            with self.subTest(period=period):
                result = self.edit(lambda wb: set_row(wb, "03_HISTORICO_MENSUAL", periodo=period))
                self.assertIn(code, self.codes(result))

    def test_quantities_negative_numeric_precision_formula(self):
        for field in ["venta_real", "pedido", "enviado"]:
            for value, code in [(-1, "CANTIDAD_NEGATIVA"), ("12", "CANTIDAD_NO_NUMERICA"), (True, "CANTIDAD_NO_NUMERICA"), (1.1234, "MAS_DE_TRES_DECIMALES"), ("=1+1", "FORMULA_NO_PERMITIDA"), (10**15, "CANTIDAD_FUERA_DE_RANGO")]:
                with self.subTest(field=field, value=value):
                    result = self.edit(lambda wb: set_row(wb, "03_HISTORICO_MENSUAL", **{field: value}))
                    self.assertIn(code, self.codes(result))

    def test_history_without_any_quantity(self):
        def remove(wb):
            for col in [4, 5, 6]:
                wb["03_HISTORICO_MENSUAL"].cell(2, col).value = None
        self.assertIn("CANTIDADES_AUSENTES", self.codes(self.edit(remove)))

    def test_exact_product_variants_and_sources(self):
        result = self.edit(lambda wb: set_row(wb, "03_HISTORICO_MENSUAL", variante_codigo="OTHER", fuente_id="MISSING"))
        self.assertIn("PRODUCTO_INEXISTENTE", self.codes(result))
        self.assertIn("FUENTE_INEXISTENTE", self.codes(result))

    def test_duplicate_rows_count_once_even_multiple_keys(self):
        def duplicate(wb):
            set_row(wb, "02_PRODUCTOS", producto_id_supabase="00000000-0000-4000-8000-000000000001")
            for sheet in ["02_PRODUCTOS", "03_HISTORICO_MENSUAL"]:
                for col in range(1, len(headers(sheet)) + 1):
                    wb[sheet].cell(3, col, wb[sheet].cell(2, col).value)
        result = self.edit(duplicate)
        self.assertEqual(result[0]["duplicados"], 2)
        self.assertEqual(sum(i["codigo"] == "DUPLICADO" for i in result[1]), 3)

    def test_forecast_future_targets_versions_and_emission(self):
        def forecast(wb):
            for row, version in [(2, 1), (3, 2)]:
                set_row(wb, "04_FORECAST_CLIENTE", row, codigo_producto_origen="000001", version_forecast_cliente=version,
                        periodo_emision="2024-12", periodo_objetivo="2025-02", cantidad_forecast_cliente=100.125,
                        fecha_recepcion=date(2024, 12, 20), estado_documento="DRAFT", fuente_id="SYNTH_SOURCE")
        result = self.edit(forecast)
        self.assertEqual(result[0]["errores"], 0)
        wb = load_workbook(self.file)
        wb["04_FORECAST_CLIENTE"]["D2"].value = None
        wb["04_FORECAST_CLIENTE"]["C3"].value = "C1"
        wb.save(self.file)
        wb.close()
        codes = self.codes(validate(self.file))
        self.assertIn("FORECAST_SIN_EMISION", codes)
        self.assertIn("VERSION_INVALIDA", codes)

    def test_pending_ambiguous_and_bata_size(self):
        result = self.edit(lambda wb: set_row(wb, "02_PRODUCTOS", modelo="UNKNOWN", color="SURTIDO", tamano_presentacion="BATAS"))
        self.assertTrue({"CLASIFICACION_PENDIENTE", "COLOR_AMBIGUO", "BATA_ES_TIPO_NO_TAMANO"}.issubset(self.codes(result)))
        result = self.edit(lambda wb: set_row(wb, "02_PRODUCTOS", color="POR DEFINIR", tamano_presentacion="POR CONFIRMAR"))
        self.assertIn("TAMANO_AMBIGUO", self.codes(result))

    def test_missing_identity_invalid_uuid_and_dates(self):
        def invalid(wb):
            wb["02_PRODUCTOS"]["A2"].value = None
            wb["02_PRODUCTOS"]["B2"] = "not-a-uuid"
            wb["01_CONTROL_CARGA"]["F2"] = "2024-02-30"
        result = self.edit(invalid)
        self.assertTrue({"CAMPO_OBLIGATORIO", "UUID_INVALIDO", "FECHA_INVALIDA", "PRODUCTO_INEXISTENTE"}.issubset(self.codes(result)))
        self.assertEqual(result[0]["productos_sin_identidad_suficiente"], 1)
        result = self.edit(lambda wb: set_row(wb, "02_PRODUCTOS", codigo_producto_origen=123))
        self.assertIn("IDENTIFICADOR_NUMERICO", self.codes(result))
        self.assertEqual(result[0]["productos_sin_identidad_suficiente"], 1)

    def test_forecast_duplicate_negative_and_version_metadata(self):
        def invalid(wb):
            for row in [2, 3]:
                set_row(wb, "04_FORECAST_CLIENTE", row, codigo_producto_origen="000001", version_forecast_cliente=1,
                        periodo_emision="2024-12", periodo_objetivo="2025-01", cantidad_forecast_cliente=-2,
                        fecha_recepcion=date(2024, 12, 19 + row), estado_documento="DRAFT", fuente_id="SYNTH_SOURCE")
        result = self.edit(invalid)
        self.assertTrue({"DUPLICADO", "CANTIDAD_NEGATIVA", "VERSION_CLIENTE_INCONSISTENTE"}.issubset(self.codes(result)))

    def test_missing_model_is_warning_not_hidden_product(self):
        def remove(wb):
            wb["02_PRODUCTOS"]["H2"].value = None
        result = self.edit(remove)
        self.assertIn("MODELO_SIN_CLASIFICACION", self.codes(result))
        self.assertEqual(result[0]["productos_registrados"], 1)

    def test_dictionary_pending_and_conflicting(self):
        def rules(wb):
            set_row(wb, "06_DICCIONARIOS", 11, atributo="COLOR", valor_original="COMBINADO", valor_homologado="AZUL",
                    estado_equivalencia="PENDIENTE", fuente_regla="Evidencia sintética")
            set_row(wb, "02_PRODUCTOS", color="COMBINADO")
            set_row(wb, "06_DICCIONARIOS", 12, atributo="TAMANO_PRESENTACION", valor_original="MB", valor_homologado="OTRO",
                    estado_equivalencia="CONFIRMADA", fuente_regla="Evidencia sintética contradictoria")
        result = self.edit(rules)
        self.assertTrue({"EQUIVALENCIA_PENDIENTE", "COLOR_AMBIGUO", "EQUIVALENCIA_CONTRADICTORIA"}.issubset(self.codes(result)))

    def test_exact_headers_missing_extra_sheets_and_data_columns(self):
        for mutate, code in [(lambda wb: setattr(wb["02_PRODUCTOS"]["A1"], "value", "Codigo_producto_origen"), "ENCABEZADOS_NO_COINCIDEN"),
                             (lambda wb: wb.remove(wb["05_FUENTES"]), "HOJA_FALTANTE"),
                             (lambda wb: wb.create_sheet("EXTRA"), "HOJAS_NO_COINCIDEN"),
                             (lambda wb: wb["02_PRODUCTOS"].cell(2, 15, "EXTRA"), "COLUMNA_EXTRA")]:
            self.file = generate(self.file)
            self.assertIn(code, self.codes(self.edit(mutate)))

    def test_source_hash_is_real_file_and_no_path_escape(self):
        self.edit()
        source = self.base / "synthetic.csv"
        source.write_bytes(b"synthetic isolated source\n")
        summary, _, status = validate(self.file, self.base)
        self.assertEqual(summary["errores"], 0)
        self.assertEqual(status["source_hashes"][0]["sha256"], sha256_file(source))
        result = self.edit(lambda wb: set_row(wb, "05_FUENTES", nombre_archivo="../outside.csv"))
        self.assertEqual(result[0]["errores"], 0)
        self.assertIn("ARCHIVO_FUENTE_NO_LOCALIZADO", self.codes(validate(self.file, self.base)))

    def test_reports_and_consultation_copy_never_overwrite_input(self):
        summary, issues, status = self.edit()
        original = sha256_file(self.file)
        report = self.base / "consultation.xlsx"
        output = self.base / "reports"
        write_reports(self.file, output, summary, issues, status, report)
        self.assertEqual(len(list(output.iterdir())), 3)
        self.assertEqual(json.loads((output / "validation_summary.json").read_text(encoding="utf-8"))["errores"], 0)
        wb = load_workbook(report)
        self.assertEqual(wb["07_VALIDACIONES"]["B4"].value, 1)
        self.assertEqual(wb["07_VALIDACIONES"]["B11"].value, 0)
        self.assertTrue(wb["02_PRODUCTOS"]["A2"].protection.locked)
        wb.close()
        self.assertEqual(original, sha256_file(self.file))
        with self.assertRaises(ValueError):
            write_reports(self.file, output, summary, issues, status, self.file)

    def test_no_network_and_read_only_input(self):
        self.edit()
        before = sha256_file(self.file)
        with patch("socket.socket", side_effect=AssertionError("Network forbidden")):
            summary, _, status = validate(self.file)
        self.assertEqual(summary["errores"], 0)
        self.assertEqual(before, sha256_file(self.file))
        self.assertFalse(status["supabase_connected"])

    def test_unknown_formulas_rejected_and_report_is_literal(self):
        result = self.edit(lambda wb: setattr(wb["07_VALIDACIONES"]["D20"], "value", '=WEBSERVICE("https://example.invalid")'))
        self.assertIn("FORMULA_NO_PERMITIDA", self.codes(result))
        copy = self.base / "rejected-consultation.xlsx"
        write_reports(self.file, self.base / "reports", *result, report_workbook=copy)
        wb = load_workbook(copy)
        self.assertEqual(wb["07_VALIDACIONES"]["D20"].data_type, "s")
        wb.close()

    def test_public_gate_rejects_business_data(self):
        self.edit()
        self.assertTrue(any("business_data_present" in p for p in audit_template(self.file)))

    def test_macro_external_link_and_corruption_rejected(self):
        for part, content, code in [("xl/vbaProject.bin", b"fake", "MACROS_NO_PERMITIDAS"),
                                    ("xl/connections.xml", b"<connections/>", "CONEXIONES_NO_PERMITIDAS"),
                                    ("xl/_rels/evil.rels", b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>', "VINCULO_EXTERNO_NO_PERMITIDO")]:
            self.file = generate(self.file)
            with zipfile.ZipFile(self.file, "a") as archive:
                archive.writestr(part, content)
            self.assertIn(code, self.codes(validate(self.file)))
            self.assertNotEqual(audit_template(self.file), [])
        self.file.write_bytes(b"not an xlsx")
        self.assertIn("ARCHIVO_NO_VALIDABLE", self.codes(validate(self.file)))

    def test_thousands_unique_products_no_synthetic_missing_months(self):
        def many(wb):
            for row in range(3, 2003):
                code = f"SYNTH_{row:06d}"
                set_row(wb, "02_PRODUCTOS", row, codigo_producto_origen=code, descripcion_original="Producto sintético",
                        modelo="MODELO NUEVO", fuente_id="SYNTH_SOURCE")
                for month in range(1, 7):
                    set_row(wb, "03_HISTORICO_MENSUAL", 3 + (row - 3) * 6 + month - 1,
                            codigo_producto_origen=code, periodo=f"2024-{month:02d}", venta_real=month * 1.125, fuente_id="SYNTH_SOURCE")
        result = self.edit(many)
        self.assertEqual(result[0]["registros_historicos"], 12001)
        self.assertEqual(result[0]["productos_registrados"], 2001)
        self.assertEqual(result[0]["errores"], 0)
        self.assertEqual(result[0]["duplicados"], 0)

    def test_cli_only_prints_summary_and_correct_exit(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts/validate_annual_template.py"), str(self.file),
                                 "--output-dir", str(self.base / "cli")], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertEqual(json.loads(result.stdout)["estado_general"], "ERROR")
        self.assertEqual(len(result.stdout.splitlines()), 1)


if __name__ == "__main__":
    unittest.main()
