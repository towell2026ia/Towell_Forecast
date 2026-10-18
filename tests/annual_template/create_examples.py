"""Rebuild local report examples with isolated synthetic data, never business data."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from generate_annual_template import generate  # noqa: E402
from validate_annual_template import validate, write_reports  # noqa: E402
from test_annual_template import synthetic  # noqa: E402


def main():
    destination = ROOT / "docs/examples/annual_template"
    with tempfile.TemporaryDirectory() as directory:
        for case in ["empty", "synthetic_warning", "synthetic_errors"]:
            source = generate(Path(directory) / f"{case}.xlsx")
            if case != "empty":
                wb = load_workbook(source)
                synthetic(wb)
                if case == "synthetic_errors":
                    for col in range(1, 9):
                        wb["03_HISTORICO_MENSUAL"].cell(3, col, wb["03_HISTORICO_MENSUAL"].cell(2, col).value)
                    wb["03_HISTORICO_MENSUAL"]["D3"] = -1
                    wb["03_HISTORICO_MENSUAL"]["G3"] = "MISSING_SYNTHETIC"
                wb.save(source)
                wb.close()
            summary, issues, status = validate(source)
            status["example_kind"] = "empty_template" if case == "empty" else "isolated_synthetic_data"
            write_reports(source, destination / case, summary, issues, status)
            print(f"{case}: {summary['estado_general']}; errors={summary['errores']}; duplicates={summary['duplicados']}")


if __name__ == "__main__":
    main()
