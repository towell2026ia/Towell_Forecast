"""Fail-closed monthly import domain, independent of forecast and web runtime.

Production ports must use private Storage and a single database transaction for
confirm.  The included memory port is for tests only, never production state.
"""

from __future__ import annotations

import copy
import csv
import hashlib
import io
import re
import uuid
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import PurePath
from typing import Any, Callable, Iterator, Protocol

from openpyxl import load_workbook

from .historical_corpus import code, number


PERIOD = re.compile(r"^20\d\d-(0[1-9]|1[0-2])$")
METRICS = {"sales": "SALES", "order": "ORDER", "delivery": "DELIVERY"}
REQUIRED = ("product_code", "description", "period")


@dataclass(frozen=True)
class ImportProfileVersion:
    profile_id: str
    chain_id: str
    version: int
    columns: dict[str, str]
    sheet: str | None = None

    def validate(self) -> None:
        if self.version < 1 or not self.chain_id or not self.profile_id:
            raise ValueError("invalid_profile_identity")
        if any(not self.columns.get(field) for field in REQUIRED):
            raise ValueError("incomplete_profile_mapping")
        if not any(self.columns.get(metric) for metric in METRICS):
            raise ValueError("profile_requires_metric")


@dataclass(frozen=True)
class ParsedRow:
    row_no: int
    product_code: str
    variant_code: str | None
    description: str
    category: str
    period: str
    metrics: dict[str, str]


@dataclass
class Batch:
    id: str
    chain_id: str
    profile: ImportProfileVersion
    filename: str
    sha256: str
    storage_path: str
    uploaded_at: str
    uploaded_by: str
    status: str = "UPLOADED"
    row_count: int = 0
    valid_rows: int = 0
    rejected_rows: int = 0
    preview: dict[str, Any] | None = None
    rows: list[ParsedRow] = field(default_factory=list)


class MonthlyImportPort(Protocol):
    def find_batch(self, chain_id: str, digest: str) -> Batch | None: ...
    def put_source(self, storage_path: str, content: bytes) -> None: ...
    def source(self, storage_path: str) -> bytes: ...
    def add_batch(self, batch: Batch) -> None: ...
    def batch(self, batch_id: str) -> Batch: ...
    def product(self, chain_id: str, product_code: str, variant_code: str | None) -> dict[str, Any] | None: ...
    def categories(self, chain_id: str) -> set[str]: ...
    def confirm_rows(self, batch: Batch, rows: list[ParsedRow]) -> None: ...


def _csv_rows(content: bytes) -> list[tuple[int, dict[str, Any], set[str]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("csv_must_be_utf8") from error
    reader = csv.DictReader(io.StringIO(text, newline=""))
    if (not reader.fieldnames or any(not field for field in reader.fieldnames)
            or len(set(reader.fieldnames)) != len(reader.fieldnames)):
        raise ValueError("invalid_csv_header")
    return [(index, dict(row), set()) for index, row in enumerate(reader, 2)]


def _xlsx_rows(content: bytes, sheet_name: str | None) -> list[tuple[int, dict[str, Any], set[str]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        expressions = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as error:
        raise ValueError("invalid_xlsx") from error
    try:
        name = sheet_name or workbook.sheetnames[0]
        if name not in workbook.sheetnames:
            raise ValueError("profile_sheet_missing")
        values = workbook[name].iter_rows(values_only=True)
        formulas = expressions[name].iter_rows(values_only=True)
        headers = [str(cell or "").strip() for cell in next(values)]
        next(formulas)
        if any(not item for item in headers) or len(set(headers)) != len(headers):
            raise ValueError("invalid_xlsx_header")
        rows = []
        for index, (row, formula_row) in enumerate(zip(values, formulas), 2):
            if not any(value is not None for value in row):
                continue
            record = dict(zip(headers, row))
            formulas_in_row = {headers[column] for column, value in enumerate(formula_row[:len(headers)])
                               if isinstance(value, str) and value.startswith("=")}
            rows.append((index, record, formulas_in_row))
        return rows
    finally:
        workbook.close()
        expressions.close()


def parse_monthly(content: bytes, filename: str, profile: ImportProfileVersion) -> tuple[list[ParsedRow], list[dict[str, Any]], int]:
    profile.validate()
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension == "csv":
        source_rows = _csv_rows(content)
    elif extension == "xlsx":
        source_rows = _xlsx_rows(content, profile.sheet)
    else:
        raise ValueError("unsupported_import_format")
    accepted: list[ParsedRow] = []
    rejected: list[dict[str, Any]] = []
    for row_no, raw, formula_fields in source_rows:
        def field(name: str) -> Any:
            column = profile.columns.get(name)
            return raw.get(column) if column else None

        errors = []
        product_code = code(field("product_code"))
        description = str(field("description") or "").strip()
        raw_period = field("period")
        if isinstance(raw_period, (date, datetime)):
            period = raw_period.strftime("%Y-%m") if raw_period.day == 1 else ""
        else:
            period = str(raw_period or "").strip()
        variant = str(field("variant_code") or "").strip() or None
        category = str(field("category") or "").strip() or "UNCLASSIFIED"
        if not product_code or not any(c.isdigit() for c in product_code):
            errors.append("invalid_product_code")
        if any(profile.columns.get(name) in formula_fields for name in
               ("product_code", "variant_code", "description", "category", "period")):
            errors.append("unverified_identity_formula")
        if not description:
            errors.append("description_required")
        if not PERIOD.fullmatch(period):
            errors.append("invalid_period")
        metrics: dict[str, str] = {}
        for source_name, metric in METRICS.items():
            column = profile.columns.get(source_name)
            if not column:
                continue
            value = field(source_name)
            if column in formula_fields or isinstance(value, str) and value.startswith("="):
                errors.append(f"unverified_formula:{source_name}")
                continue
            status, parsed = number(value)
            if status == "INVALID":
                errors.append(f"invalid_metric:{source_name}")
            elif status == "VALUE" and parsed is not None:
                metrics[metric] = parsed
        if not metrics:
            errors.append("no_valid_metric")
        if errors:
            rejected.append({"row": row_no, "errors": errors})
        else:
            accepted.append(ParsedRow(row_no, product_code, variant, description,
                                      category, period, metrics))
    return accepted, rejected, len(source_rows)


class MonthlyImportWorkflow:
    def __init__(self, port: MonthlyImportPort, clock: Callable[[], datetime] | None = None):
        self.port = port
        self.clock = clock or (lambda: datetime.now(timezone.utc))

    def upload(self, content: bytes, *, filename: str, profile: ImportProfileVersion,
               uploaded_by: str) -> Batch:
        profile.validate()
        if not content or not uploaded_by or not filename:
            raise ValueError("invalid_upload")
        basename = PurePath(filename.replace("\\", "/")).name
        if basename != filename or basename in {".", ".."}:
            raise ValueError("unsafe_filename")
        digest = hashlib.sha256(content).hexdigest()
        existing = self.port.find_batch(profile.chain_id, digest)
        if existing:
            return existing
        timestamp = self.clock()
        if timestamp.tzinfo is None:
            raise ValueError("server_clock_must_be_aware")
        uploaded_at = timestamp.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
        storage_path = f"{profile.chain_id}/{digest}/{basename}"
        # The production port must write this to the private source-files bucket.
        self.port.put_source(storage_path, content)
        batch = Batch(str(uuid.uuid4()), profile.chain_id, profile, basename, digest,
                      storage_path, uploaded_at, uploaded_by)
        self.port.add_batch(batch)
        return batch

    def validate(self, batch_id: str) -> dict[str, Any]:
        batch = self.port.batch(batch_id)
        if batch.status not in {"UPLOADED", "VALIDATED"}:
            raise ValueError("invalid_batch_state")
        content = self.port.source(batch.storage_path)
        if hashlib.sha256(content).hexdigest() != batch.sha256:
            raise ValueError("source_hash_mismatch")
        rows, rejected, total = parse_monthly(content, batch.filename, batch.profile)
        existing = new = updated = 0
        new_categories = set()
        for row in rows:
            product = self.port.product(batch.chain_id, row.product_code, row.variant_code)
            if product is None:
                new += 1
            else:
                existing += 1
                if product["description"] != row.description:
                    updated += 1
            if row.category not in self.port.categories(batch.chain_id):
                new_categories.add(row.category)
        periods = sorted({row.period for row in rows})
        preview = {"chain_id": batch.chain_id, "periods": periods, "rows": total,
                   "valid_rows": len(rows), "rejected_rows": len(rejected),
                   "products_existing": existing, "products_new": new,
                   "products_updated": updated, "categories_new": len(new_categories),
                   "warnings": [], "blocking_errors": rejected}
        batch.rows = rows
        batch.row_count = total
        batch.valid_rows = len(rows)
        batch.rejected_rows = len(rejected)
        batch.preview = preview
        batch.status = "VALIDATED" if rows and not rejected else "REJECTED"
        return copy.deepcopy(preview)

    def preview(self, batch_id: str) -> dict[str, Any]:
        batch = self.port.batch(batch_id)
        if batch.preview is None:
            raise ValueError("validation_required")
        return copy.deepcopy(batch.preview)

    def confirm(self, batch_id: str, *, actor_id: str) -> Batch:
        batch = self.port.batch(batch_id)
        if batch.status != "VALIDATED" or batch.preview is None or actor_id != batch.uploaded_by:
            raise ValueError("confirmation_forbidden")
        if len(batch.rows) != batch.valid_rows or batch.rejected_rows:
            raise ValueError("unreconciled_batch")
        # The production implementation of this port MUST execute all row
        # changes and final batch transition inside one PostgreSQL transaction.
        self.port.confirm_rows(batch, batch.rows)
        return self.port.batch(batch_id)


class MemoryImportPort:
    """Deterministic acceptance-test port; not a production persistence adapter."""

    def __init__(self):
        self.batches: dict[str, Batch] = {}
        self.files: dict[str, bytes] = {}
        self.product_rows: dict[tuple[str, str, str | None], dict[str, Any]] = {}
        self.category_rows: set[tuple[str, str]] = set()
        self.observations: list[dict[str, Any]] = []
        self.audit: list[dict[str, Any]] = []
        self.fail_after: int | None = None

    def find_batch(self, chain_id: str, digest: str) -> Batch | None:
        return next((batch for batch in self.batches.values()
                     if batch.chain_id == chain_id and batch.sha256 == digest), None)

    def put_source(self, storage_path: str, content: bytes) -> None:
        self.files[storage_path] = content

    def source(self, storage_path: str) -> bytes:
        return self.files[storage_path]

    def add_batch(self, batch: Batch) -> None:
        self.batches[batch.id] = batch

    def batch(self, batch_id: str) -> Batch:
        return self.batches[batch_id]

    def product(self, chain_id: str, product_code: str, variant_code: str | None) -> dict[str, Any] | None:
        return self.product_rows.get((chain_id, product_code, variant_code))

    def categories(self, chain_id: str) -> set[str]:
        return {name for chain, name in self.category_rows if chain == chain_id}

    @contextmanager
    def transaction(self) -> Iterator[None]:
        snapshot = copy.deepcopy((self.batches, self.product_rows, self.category_rows,
                                  self.observations, self.audit))
        try:
            yield
        except Exception:
            (self.batches, self.product_rows, self.category_rows,
             self.observations, self.audit) = snapshot
            raise

    def confirm_rows(self, batch: Batch, rows: list[ParsedRow]) -> None:
        with self.transaction():
            count = 0
            for row in rows:
                key = (batch.chain_id, row.product_code, row.variant_code)
                self.category_rows.add((batch.chain_id, row.category))
                product = self.product_rows.get(key)
                if product is None:
                    product = {"id": str(uuid.uuid4()), "chain_id": batch.chain_id,
                               "product_code": row.product_code, "variant_code": row.variant_code,
                               "description": row.description, "category": row.category,
                               "status": "NEW", "first_seen_period": row.period,
                               "last_seen_period": row.period, "created_from_batch_id": batch.id}
                    self.product_rows[key] = product
                else:
                    product["description"] = row.description
                    product["category"] = row.category
                    product["first_seen_period"] = min(product["first_seen_period"], row.period)
                    product["last_seen_period"] = max(product["last_seen_period"], row.period)
                for metric, value in row.metrics.items():
                    prior = [observation for observation in self.observations
                             if observation["chain_id"] == batch.chain_id
                             and observation["product_id"] == product["id"]
                             and observation["period"] == row.period
                             and observation["metric_code"] == metric]
                    self.observations.append({"id": str(uuid.uuid4()), "chain_id": batch.chain_id,
                                              "product_id": product["id"], "period": row.period,
                                              "metric_code": metric, "value": value,
                                              "version_no": len(prior) + 1,
                                              "source_batch_id": batch.id,
                                              "available_at": batch.uploaded_at,
                                              "availability_source": "SYSTEM_INGESTION"})
                count += 1
                if self.fail_after is not None and count >= self.fail_after:
                    raise RuntimeError("injected_transaction_failure")
            self.audit.append({"batch_id": batch.id, "actor": batch.uploaded_by,
                               "event": "MONTHLY_IMPORT_CONFIRMED"})
            self.batches[batch.id].status = "IMPORTED"
