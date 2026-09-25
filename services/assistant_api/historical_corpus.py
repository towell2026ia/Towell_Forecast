"""Read-only, source-grounded scan of historical product observations.

This module never writes to a database and deliberately does not pick a winner
when snapshots disagree.  Its output is suitable for an import approval gate.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ERRORS = {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
HEADER_PERIOD = re.compile(r"(20\d{2})/(0[1-9]|1[0-2])\s+(.+)")
MONTHLY_METRICS = {
    "POS Qty": "SALES",
    "Total Eaches Str Ordered": "ORDER",
    "Total Eaches Str Received": "DELIVERY",
}


@dataclass(frozen=True)
class SourceSpec:
    path: Path
    chain_hint: str | None = None
    summary_year: int | None = None
    source_cutoff: str | None = None


@dataclass(frozen=True)
class Candidate:
    chain: str
    item: str
    upc: str
    description: str
    category: str
    period: str
    metric: str
    value: str
    source_sha256: str
    sheet: str
    cell: str
    source_cutoff: str
    formula: bool = False


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def code(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)) and value == int(value):
        return str(int(value))
    text = str(value).strip()
    return text if text and text.casefold() not in {"none", "nan", "item", "total", "-"} else ""


def number(value: Any) -> tuple[str, str | None]:
    if value is None or str(value).strip() == "":
        return "MISSING", None
    if isinstance(value, bool) or (isinstance(value, str) and value.strip().upper() in ERRORS):
        return "INVALID", None
    source = str(value).strip()
    if "," in source and not re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", source):
        return "INVALID", None
    try:
        result = Decimal(source.replace(",", ""))
    except (InvalidOperation, ValueError):
        return "INVALID", None
    if not result.is_finite() or result < 0:
        return "INVALID", None
    return "VALUE", format(result.normalize(), "f")


def _metadata(path: Path) -> dict[str, Any]:
    return {"source_file": path.name, "sha256": sha256_file(path),
            "size_bytes": path.stat().st_size, "sheets": [], "detected_fields": [],
            "source_periods": [], "chains": [], "products_detected": 0,
            "quality_issues": [], "duplicate_relationship": None}


def _wide_rows(values: Any, formulas: Any, *, chain: str, digest: str,
               sheet: str, header_row: int, aliases: dict[tuple[str, str], str],
               ambiguous_aliases: set[tuple[str, str]], issues: Counter[str]) -> list[Candidate]:
    headers = next(values.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    metric_cols: list[tuple[int, str, str]] = []
    for index, header in enumerate(headers):
        match = HEADER_PERIOD.search(str(header or ""))
        if not match:
            continue
        year, month, label = match.groups()
        if int(year) < 2023 or int(year) > 2026:
            continue
        for fragment, metric in MONTHLY_METRICS.items():
            if fragment in label:
                metric_cols.append((index, f"{year}-{month}", metric))
                break
    if not metric_cols:
        return []
    cutoff = max(period for _, period, _ in metric_cols)
    found: list[Candidate] = []
    value_rows = values.iter_rows(min_row=header_row + 1, values_only=True)
    formula_rows = formulas.iter_rows(min_row=header_row + 1, values_only=True)
    for row_no, (row, expression) in enumerate(zip(value_rows, formula_rows), header_row + 1):
        item = code(row[0] if row else None)
        if not item or not any(char.isdigit() for char in item):
            continue
        if (chain, item) in ambiguous_aliases:
            issues["unresolved_identity_rows"] += 1
            continue
        upc = aliases.get((chain, item), "")
        description = str(row[1] or "").strip()
        for index, period, metric in metric_cols:
            raw = row[index] if index < len(row) else None
            expression_value = expression[index] if index < len(expression) else None
            if isinstance(expression_value, str) and expression_value.startswith("="):
                issues["formula_without_direct_source"] += 1
                continue
            status, value = number(raw)
            if status != "VALUE":
                issues["missing" if status == "MISSING" else "invalid_value"] += 1
                continue
            found.append(Candidate(chain, item, upc, description, "UNCLASSIFIED", period,
                                   metric, value or "", digest, sheet,
                                   f"{get_column_letter(index + 1)}{row_no}", cutoff))
    return found


def _summary_rows(values: Any, formulas: Any, *, chain: str, year: int,
                  digest: str, sheet: str, issues: Counter[str],
                  source_cutoff: str) -> tuple[list[Candidate], dict[tuple[str, str], str]]:
    """The ITEM/UPC summary has literal historical cells and formula-derived cells.

    Formula-derived values are intentionally excluded; a direct Base sheet can
    supply them without trusting a stale IFERROR/XLOOKUP cache.
    """
    result: list[Candidate] = []
    aliases: dict[tuple[str, str], str] = {}
    category = "UNCLASSIFIED"
    groups = ((9, "SALES"), (23, "ORDER"), (36, "DELIVERY"))
    rows = values.iter_rows(min_row=3, max_col=49, values_only=True)
    expressions = formulas.iter_rows(min_row=3, max_col=49, values_only=True)
    for row_no, (row, expression) in enumerate(zip(rows, expressions), 3):
        item = code(row[1])
        upc = code(row[2])
        if row_no == 3 and item:
            category = item.upper()
        if not upc and item and not any(char.isdigit() for char in item):
            if item.casefold() in {"baño", "bano", "cocina"}:
                category = item.upper()
            continue
        if not item or not upc or not any(char.isdigit() for char in item):
            continue
        existing = aliases.get((chain, item))
        if existing and existing != upc:
            issues["alias_conflict"] += 1
            continue
        aliases[(chain, item)] = upc
        description = str(row[3] or "").strip()
        for start, metric in groups:
            for month in range(1, 13):
                if f"{year}-{month:02d}" > source_cutoff:
                    continue
                index = start + month - 1
                formula = expression[index]
                if isinstance(formula, str) and formula.startswith("="):
                    issues["formula_without_direct_source"] += 1
                    continue
                status, value = number(row[index])
                if status != "VALUE":
                    issues["missing" if status == "MISSING" else "invalid_value"] += 1
                    continue
                result.append(Candidate(chain, item, upc, description, category,
                                        f"{year}-{month:02d}", metric, value or "", digest,
                                        sheet, f"{get_column_letter(index + 1)}{row_no}",
                                        source_cutoff))
    return result, aliases


def _master_rows(values: Any, formulas: Any, *, digest: str, sheet: str,
                 issues: Counter[str]) -> list[Candidate]:
    result: list[Candidate] = []
    rows = values.iter_rows(min_row=18, max_col=24, values_only=True)
    expressions = formulas.iter_rows(min_row=18, max_col=24, values_only=True)
    for row_no, (row, expression) in enumerate(zip(rows, expressions), 18):
        chain, item, upc = str(row[1] or "").strip(), code(row[6]), code(row[7])
        period_value = row[16]
        if not chain or not (item or upc) or not isinstance(period_value, (date, datetime)):
            continue
        period = period_value.strftime("%Y-%m")
        if not ("2023-01" <= period <= "2026-12"):
            continue
        description = str(row[8] or "").strip()
        category = str(row[4] or "").strip() or "UNCLASSIFIED"
        for index, metric in ((21, "SALES"), (22, "ORDER"), (23, "DELIVERY")):
            formula = expression[index]
            if isinstance(formula, str) and formula.startswith("="):
                issues["formula_without_direct_source"] += 1
                continue
            status, value = number(row[index])
            if status != "VALUE":
                issues["missing" if status == "MISSING" else "invalid_value"] += 1
                continue
            result.append(Candidate(chain, item, upc, description, category, period,
                                    metric, value or "", digest, sheet,
                                    f"{get_column_letter(index + 1)}{row_no}", "2026-07"))
    cutoff = max((row.period for row in result), default="")
    return [replace(row, source_cutoff=cutoff) for row in result]


def scan_sources(sources: Iterable[SourceSpec]) -> dict[str, Any]:
    """Inventory and reconcile without modifying source files or remote state."""
    specs = list(sources)
    for spec in specs:
        if spec.summary_year is not None and (not spec.chain_hint or not 2023 <= spec.summary_year <= 2026):
            raise ValueError("invalid_summary_profile")
        if spec.source_cutoff and (not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", spec.source_cutoff)
                                   or spec.summary_year and not spec.source_cutoff.startswith(str(spec.summary_year))):
            raise ValueError("invalid_source_cutoff")
    manifests: list[dict[str, Any]] = []
    candidates: list[Candidate] = []
    global_issues: Counter[str] = Counter()
    known_hashes: dict[str, str] = {}
    alias_pairs: dict[tuple[str, str], set[str]] = defaultdict(set)
    # Pass one: record explicit UPC/ITEM pairs; never infer an alias from text.
    for spec in specs:
        path = spec.path.resolve(strict=True)
        manifest = _metadata(path)
        if manifest["sha256"] in known_hashes:
            manifest["duplicate_relationship"] = f"exact_file:{known_hashes[manifest['sha256']]}"
            global_issues["exact_file_duplicate"] += 1
        else:
            known_hashes[manifest["sha256"]] = path.name
        manifests.append(manifest)
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook:
            manifest["sheets"].append({"sheet": sheet.title, "row_count": sheet.max_row,
                                        "max_column": sheet.max_column,
                                        "classification": "UNSUPPORTED_OR_AGGREGATE",
                                        "source_periods": [], "chains": [],
                                        "products_detected": 0, "parsed_observations": 0})
            if spec.chain_hint and spec.summary_year and sheet.title.endswith("(2)"):
                header = next(sheet.iter_rows(min_row=4, max_row=4, max_col=4, values_only=True))
                if tuple(str(x or "").strip() for x in header[1:4]) == ("ITEM", "UPC", "Modelo"):
                    for row in sheet.iter_rows(min_row=5, max_col=4, values_only=True):
                        item, upc = code(row[1]), code(row[2])
                        if item and upc and any(c.isdigit() for c in item):
                            alias_pairs[(spec.chain_hint, item)].add(upc)
            if sheet.title == "BAASE":
                header = next(sheet.iter_rows(min_row=17, max_row=17, max_col=24, values_only=True))
                if str(header[1]).strip() == "Cadena" and str(header[16]).strip() == "Fecha":
                    for row in sheet.iter_rows(min_row=18, max_col=8, values_only=True):
                        chain, item, upc = str(row[1] or "").strip(), code(row[6]), code(row[7])
                        if chain and item and upc and any(c.isdigit() for c in item):
                            alias_pairs[(chain, item)].add(upc)
        workbook.close()
    ambiguous_aliases = {key for key, upcs in alias_pairs.items() if len(upcs) > 1}
    alias_map = {key: next(iter(upcs)) for key, upcs in alias_pairs.items() if len(upcs) == 1}
    # Pass two: direct source cells only; ignore aggregate and forecast sheets.
    for spec, manifest in zip(specs, manifests):
        if manifest["duplicate_relationship"]:
            manifest["quality_issues"] = {"exact_file_duplicate_skipped": 1}
            manifest["parsed_observations"] = 0
            manifest["date_range"] = None
            for sheet_record in manifest["sheets"]:
                sheet_record["classification"] = "DUPLICATE_FILE_SKIPPED"
            continue
        workbook = load_workbook(spec.path, read_only=True, data_only=True)
        formulas = load_workbook(spec.path, read_only=True, data_only=False)
        sheet_issues: Counter[str] = Counter()
        before = len(candidates)
        for sheet in workbook:
            expression = formulas[sheet.title]
            sheet_start = len(candidates)
            sheet_record = next(record for record in manifest["sheets"] if record["sheet"] == sheet.title)
            if sheet.title == "BAASE":
                header = next(sheet.iter_rows(min_row=17, max_row=17, max_col=24, values_only=True))
                if str(header[1]).strip() == "Cadena" and str(header[16]).strip() == "Fecha":
                    candidates.extend(_master_rows(sheet, expression, digest=manifest["sha256"],
                                                   sheet=sheet.title, issues=sheet_issues))
                    manifest["detected_fields"].append("BAASE: Cadena/ITEM/UPC/Fecha/Venta/Pedido/Entrega")
                    sheet_record["classification"] = "PRODUCT_ACTUAL_MASTER"
            elif spec.chain_hint and sheet.title == "BD (2)" and spec.summary_year:
                rows, _ = _summary_rows(sheet, expression, chain=spec.chain_hint,
                                        year=spec.summary_year, digest=manifest["sha256"],
                                        sheet=sheet.title, issues=sheet_issues,
                                        source_cutoff=spec.source_cutoff or f"{spec.summary_year}-12")
                candidates.extend(rows)
                manifest["detected_fields"].append("BD (2): ITEM/UPC/Venta/Pedido/Entrega")
                sheet_record["classification"] = "PRODUCT_ACTUAL_SUMMARY_LITERALS_ONLY"
            elif spec.chain_hint and sheet.title in {"Base", "BD WM 2024"}:
                header_row = 1 if sheet.title == "Base" else 4
                header = next(sheet.iter_rows(min_row=header_row, max_row=header_row,
                                              max_col=2, values_only=True))
                if str(header[0]).strip() == "Prime Item Nbr":
                    candidates.extend(_wide_rows(sheet, expression, chain=spec.chain_hint,
                                                 digest=manifest["sha256"], sheet=sheet.title,
                                                 header_row=header_row, aliases=alias_map,
                                                 ambiguous_aliases=ambiguous_aliases,
                                                 issues=sheet_issues))
                    manifest["detected_fields"].append(f"{sheet.title}: Prime Item Nbr/POS Qty/Ordered/Received")
                    sheet_record["classification"] = "PRODUCT_ACTUAL_WIDE"
            sheet_rows = candidates[sheet_start:]
            sheet_record["source_periods"] = sorted({row.period for row in sheet_rows})
            sheet_record["chains"] = sorted({row.chain for row in sheet_rows})
            sheet_record["products_detected"] = len({(row.chain, row.upc or row.item)
                                                      for row in sheet_rows})
            sheet_record["parsed_observations"] = len(sheet_rows)
        included = candidates[before:]
        manifest["source_periods"] = sorted({r.period for r in included})
        manifest["date_range"] = [min(manifest["source_periods"]), max(manifest["source_periods"])] if included else None
        manifest["chains"] = sorted({r.chain for r in included})
        manifest["products_detected"] = len({(r.chain, r.upc or r.item) for r in included})
        manifest["parsed_observations"] = len(included)
        manifest["quality_issues"] = dict(sheet_issues)
        global_issues.update(sheet_issues)
        if not included:
            manifest["quality_issues"]["no_product_actual_profile"] = 1
            global_issues["no_product_actual_profile"] += 1
        workbook.close()
        formulas.close()
    global_issues["alias_conflict"] = len(ambiguous_aliases)
    return reconcile(manifests, candidates, alias_map, global_issues)


def reconcile(manifests: list[dict[str, Any]], candidates: list[Candidate],
              aliases: dict[tuple[str, str], str], issues: Counter[str]) -> dict[str, Any]:
    # A code may be an ITEM in one workbook and UPC in another. An explicit
    # source pair is the only allowed bridge; descriptions never form identity.
    by_product: dict[tuple[str, str], list[Candidate]] = defaultdict(list)
    for row in candidates:
        canonical = row.upc or aliases.get((row.chain, row.item)) or row.item
        by_product[(row.chain, canonical)].append(row)
    first_active: dict[tuple[str, str], str | None] = {}
    for product, rows in by_product.items():
        positive = [r.period for r in rows if Decimal(r.value) > 0]
        first_active[product] = min(positive) if positive else None
    grouped: dict[tuple[str, str, str, str], list[Candidate]] = defaultdict(list)
    not_active = unknown_zero = 0
    for product, rows in by_product.items():
        start = first_active[product]
        for row in rows:
            if Decimal(row.value) == 0 and (start is None or row.period < start):
                if start is None:
                    unknown_zero += 1
                else:
                    not_active += 1
                continue
            grouped[(product[0], product[1], row.period, row.metric)].append(row)
    equivalent = 0
    conflicts: list[dict[str, Any]] = []
    selected: list[dict[str, Any]] = []
    for key, rows in sorted(grouped.items()):
        distinct = {r.value for r in rows}
        if len(distinct) > 1:
            conflicts.append({"key": key, "values": sorted(distinct),
                              "sources": [dict(source_sha256=r.source_sha256,
                                               sheet=r.sheet, cell=r.cell, value=r.value)
                                          for r in rows]})
            continue
        if len(rows) > 1:
            equivalent += len(rows) - 1
        # Selection among equivalent values is deterministic and uses parsed
        # cutoff plus source integrity, not filename. Conflicts have no winner.
        winner = max(rows, key=lambda r: (r.source_cutoff, not r.formula,
                                          r.source_sha256, r.sheet, r.cell))
        selected.append({"key": key, "value": winner.value, "source_sha256": winner.source_sha256,
                         "sheet": winner.sheet, "cell": winner.cell,
                         "equivalent_sources": [dict(source_sha256=r.source_sha256,
                                                     sheet=r.sheet, cell=r.cell)
                                                for r in rows],
                         "availability_source": "UNKNOWN", "available_at": None})
    metric_counts = Counter(row["key"][3] for row in selected)
    chains = Counter(product[0] for product in by_product)
    periods = sorted({row["key"][2] for row in selected})
    new_by_year = Counter(start[:4] for start in first_active.values() if start)
    descriptions = defaultdict(set)
    categories = defaultdict(set)
    for product, rows in by_product.items():
        descriptions[product].update(r.description for r in rows if r.description)
        categories[product[0]].update(r.category for r in rows if r.category != "UNCLASSIFIED")
    report = {
        "manifest": manifests,
        "summary": {
            "files": len(manifests), "chains": dict(sorted(chains.items())),
            "products_unique": len(by_product), "period_min": periods[0] if periods else None,
            "period_max": periods[-1] if periods else None,
            "categories_by_chain": {chain: len(names) for chain, names in sorted(categories.items())},
            "observations": dict(sorted(metric_counts.items())),
            "confirmed_zero": sum(row["value"] == "0" for row in selected),
            "missing": issues["missing"], "not_active": not_active,
            "zero_without_activation_evidence": unknown_zero,
            "duplicate_equivalent_sources": equivalent,
            "conflicts": len(conflicts), "formula_without_direct_source": issues["formula_without_direct_source"],
            "aliases_item_to_upc": len(aliases),
            "description_changes": sum(len(names) > 1 for names in descriptions.values()),
            "products_new_by_year": dict(sorted(new_by_year.items())),
            "unknown_temporal_records": len(selected), "evidence_temporal_records": 0,
        },
        "quality_issues": dict(issues), "conflicts": conflicts, "selected": selected,
        "blocking": bool(conflicts or issues["alias_conflict"] or unknown_zero or not periods
                         or issues["formula_without_direct_source"]
                         or issues["no_product_actual_profile"]),
    }
    return report
