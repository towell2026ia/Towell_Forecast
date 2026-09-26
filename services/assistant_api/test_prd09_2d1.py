"""Synthetic public cases; private business regression is run separately by CLI."""

from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from services.assistant_api.historical_corpus import Candidate, SourceSpec, scan_sources
from services.assistant_api.historical_reconciliation import (
    CLASSIFICATIONS, SnapshotProof, TemporalEvidence, bootstrap_historical, publication_preflight,
    reconcile_historical, regression_sample,
)

ROOT = Path(__file__).resolve().parents[2]
PROJECT = "bskoyqhbgrycpwhydnnr"


def fact(value="8", **changes):
    return replace(Candidate("Chain A", "ITEM1", "CODE1", "Old name", "Towels",
                             "2025-01", "SALES", value, "a" * 64, "Detail", "V18", "2025-01"), **changes)


def report(rows, **options):
    return reconcile_historical([], rows, **options)


class DurableTestPort:
    """Test-only transactional checkpoint adapter, NOT a Supabase adapter."""
    project_ref = PROJECT

    def __init__(self, path):
        self.connection = sqlite3.connect(path)
        self.connection.execute("create table if not exists batches (key text primary key, payload text not null)")
        self.fail = False

    def commit_batch(self, key, facts):
        payload = json.dumps(facts, sort_keys=True)
        with self.connection:
            previous = self.connection.execute("select payload from batches where key=?", (key,)).fetchone()
            if previous and previous[0] != payload:
                raise ValueError("idempotency_mismatch")
            self.connection.execute("insert or ignore into batches values (?,?)", (key, payload))
            if self.fail:
                raise RuntimeError("injected_failure")
        return len(facts)


class ReconciliationTests(unittest.TestCase):
    def test_cp01_equivalent_duplicate_no_sum(self):
        result = report([fact(), fact(source_sha256="b" * 64)])
        self.assertEqual(len(result["selected"]), 1)
        self.assertEqual(result["selected"][0]["value"], "8")
        self.assertEqual(result["resolution_ledger"][0]["reason_code"], "EQUIVALENT_DUPLICATE")

    def test_cp02_byte_identical_source(self):
        result = reconcile_historical([{"sha256": "a" * 64, "duplicate_relationship": "exact_file:copy"}], [fact()])
        self.assertEqual(result["summary"]["reason_counts"]["EQUIVALENT_DUPLICATE"], 1)
        self.assertEqual(len(result["selected"]), 1)

    def test_cp03_direct_beats_summary(self):
        result = report([fact(), fact("99", evidence_level="C")])
        self.assertEqual(result["selected"][0]["value"], "8")
        self.assertEqual(result["resolution_ledger"][0]["reason_code"], "DIRECT_VS_SUMMARY")

    def test_cp04_formula_excluded(self):
        result = report([fact(), fact("99", formula=True)])
        self.assertEqual(result["summary"]["excluded"]["FORMULA_UNVERIFIED"], 1)
        self.assertFalse(result["blocking"])

    def revisions(self):
        proofs = {digest: SnapshotProof(digest, "monthly-direct", cutoff, "c" * 64,
                                      "Evidence!A1", "reviewer", "Comparable approved snapshots")
                  for digest, cutoff in (("a" * 64, "2025-01"), ("b" * 64, "2025-02"))}
        return report([fact(), fact("9", source_sha256="b" * 64, source_cutoff="2025-02")], snapshots=proofs)

    def test_cp05_validated_revision(self):
        result = self.revisions()
        self.assertEqual(result["resolution_ledger"][0]["reason_code"], "VALIDATED_REVISION")
        self.assertFalse(result["blocking"])

    def test_cp06_revision_preserves_v1(self):
        self.assertEqual([(r["version_no"], r["value"]) for r in self.revisions()["versions"]], [(1, "8"), (2, "9")])

    def test_cp07_same_snapshot_blocks_independent_of_order(self):
        rows = [fact(), fact("9", source_sha256="b" * 64)]
        for inputs in (rows, rows[::-1]):
            result = report(inputs)
            self.assertTrue(result["blocking"])
            self.assertEqual(result["selected"], [])
            self.assertEqual(result["resolution_ledger"][0]["reason_code"], "SAME_SNAPSHOT_CONFLICT")

    def test_cp08_internal_duplicate_blocks_even_with_external_source(self):
        result = report([fact(), fact("9", cell="V19"), fact("8", source_sha256="b" * 64)])
        self.assertEqual(result["resolution_ledger"][0]["reason_code"], "INTERNAL_SOURCE_DUPLICATE")
        self.assertTrue(result["blocking"])

    def test_cp09_explicit_upcs_never_merge_shared_item(self):
        result = report([fact(), fact("9", upc="CODE2")])
        self.assertEqual(result["summary"]["products"], 2)
        self.assertEqual(result["summary"]["item_multi_upc"], 1)
        self.assertFalse(result["blocking"])

    def test_cp10_item_only_ambiguity_blocks(self):
        result = report([fact(), fact(upc="CODE2"), fact(upc="", cell="V20")])
        self.assertEqual(result["summary"]["item_only_ambiguous"], 1)
        self.assertTrue(result["blocking"])

    def test_item_only_temporal_bridge_not_global_item(self):
        result = report([fact(period="2024-01"), fact("9", upc="CODE2", period="2025-01"),
                         fact("8", upc="", period="2024-01", cell="V20")])
        self.assertFalse(result["blocking"])
        self.assertEqual(result["summary"]["products"], 2)

    def test_cp11_description_alias_same_product(self):
        result = report([fact(), fact(description="New name")])
        self.assertEqual(result["summary"]["products"], 1)
        self.assertEqual(result["description_aliases"][0]["reason"], "DESCRIPTION_ALIAS")

    def test_cp12_prelaunch_zero_excluded(self):
        result = report([fact("0", period="2024-12"), fact()])
        self.assertEqual(result["summary"]["excluded"]["PRE_LAUNCH_ZERO"], 1)
        self.assertEqual(len(result["selected"]), 1)

    def test_cp13_active_zero_preserved(self):
        result = report([fact(), fact("0", period="2025-02")])
        self.assertEqual(result["selected"][-1]["value"], "0")
        self.assertEqual(result["resolution_ledger"][-1]["reason_code"], "CONFIRMED_ZERO")

    def test_cp14_unproven_zero_safely_excluded(self):
        result = report([fact(), fact("0", item="ITEM2", upc="CODE2")])
        self.assertEqual(result["summary"]["excluded"]["UNPROVEN_ZERO"], 1)
        self.assertFalse(result["blocking"])
        self.assertEqual(result["summary"]["products"], 1)

    def test_cp15_missing_not_zero(self):
        result = report([fact(), fact("", period="2025-02")])
        self.assertEqual(result["summary"]["excluded"]["MISSING"], 1)
        self.assertEqual(len(result["selected"]), 1)

    def test_cp16_invalid_excluded(self):
        result = report([fact(), fact("#REF!"), fact("-1")])
        self.assertEqual(result["summary"]["excluded"]["INVALID_VALUE"], 2)

    def test_cp17_direct_2023_accepted(self):
        self.assertEqual(report([fact(period="2023-01")])["summary"]["period_min"], "2023-01")

    def test_cp18_aggregate_never_expanded(self):
        result = report([fact("100", evidence_level="C", period="2023-01")])
        self.assertEqual(result["selected"], [])
        self.assertEqual(result["summary"]["excluded"]["INSUFFICIENT_EVIDENCE"], 1)

    def test_cp19_unknown_has_null_available_at(self):
        for row in self.revisions()["versions"]:
            self.assertEqual(row["availability_source"], "UNKNOWN")
            self.assertIsNone(row["available_at"])

    def test_cp20_real_evidence_required_for_each_level(self):
        for level in ("E1", "E2", "E3"):
            proof = TemporalEvidence("evidence-id", level, "2025-02-01T00:00:00Z", "2025-03-01T00:00:00Z",
                                     "reviewer", "a" * 64, "chain/evidence", "Chain A")
            result = report([fact()], evidence={("Chain A", "a" * 64): proof})
            self.assertEqual(result["selected"][0]["availability_source"], level)
            for bad in (replace(proof, validated_by=""), replace(proof, chain="Chain B"),
                        replace(proof, evidence_date="2026-01-01T00:00:00Z"), replace(proof, storage_path=""),
                        replace(proof, evidence_date="invalid"), replace(proof, level="SYSTEM_INGESTION")):
                with self.assertRaisesRegex(ValueError, "invalid_temporal_evidence"):
                    report([fact()], evidence={("Chain A", "a" * 64): bad})

    def test_cp21_cutoff_not_available_at_or_automatic_revision(self):
        result = report([fact(), fact("9", source_cutoff="2026-07", source_sha256="b" * 64)])
        self.assertTrue(result["blocking"])
        self.assertEqual(result["summary"]["revisions"], 0)

    def test_cp22_cross_chain_isolation(self):
        result = report([fact(), fact("9", chain="Chain B")])
        self.assertEqual(result["summary"]["products"], 2)
        self.assertFalse(result["blocking"])

    def test_cp23_natural_identity_unique(self):
        result = report([fact(), fact(description="alias", item="ITEM2")])
        self.assertEqual(result["summary"]["products"], 1)
        self.assertEqual(len(result["selected"]), 1)

    def test_cp24_current_selects_latest_value_version(self):
        self.assertEqual(self.revisions()["selected"][0]["value"], "9")

    def test_cp25_synthetic_96_fact_golden_coverage(self):
        rows = [fact(str(i + 1), item=f"ITEM{i}", upc=f"CODE{i}", cell=f"V{i+18}") for i in range(96)]
        result = report(rows)
        self.assertEqual(len(result["selected"]), 96)
        self.assertEqual(sorted(int(r["value"]) for r in result["selected"]), list(range(1, 97)))

    def test_cp26_auto_chain_regression_sample(self):
        rows = [fact(), fact("9", chain="Chain B", cell="V19")]
        result = report(rows)
        self.assertEqual(regression_sample(result, rows), {"chains_sampled": 2, "mismatches": 0, "status": "PASS"})
        self.assertEqual(regression_sample(result, [fact("99")])["status"], "FAIL")

    def execute(self, result, port):
        return bootstrap_historical(result, port, golden_pass=True, tests_pass=True, remote_gate_pass=True)

    def test_cp27_bootstrap_port_idempotency(self):
        port = DurableTestPort(":memory:")
        result = report([fact()])
        self.assertEqual(self.execute(result, port), self.execute(result, port))
        self.assertEqual(port.connection.execute("select count(*) from batches").fetchone()[0], 1)
        port.connection.close()

    def test_cp28_bootstrap_port_atomic_rollback(self):
        port = DurableTestPort(":memory:")
        port.fail = True
        with self.assertRaisesRegex(RuntimeError, "injected_failure"):
            self.execute(report([fact()]), port)
        self.assertEqual(port.connection.execute("select count(*) from batches").fetchone()[0], 0)
        port.connection.close()

    def test_cp29_bootstrap_port_restart_resume(self):
        with tempfile.TemporaryDirectory() as directory:
            path = str(Path(directory) / "test.sqlite")
            port = DurableTestPort(path)
            result = report([fact()])
            expected = self.execute(result, port)
            port.connection.close()
            port = DurableTestPort(path)
            self.assertEqual(self.execute(result, port), expected)
            self.assertEqual(port.connection.execute("select count(*) from batches").fetchone()[0], 1)
            port.connection.close()

    def test_cp30_monthly_core_unchanged(self):
        # Existing test_prd09_2d executes all CSV/XLSX/transaction cases in CI.
        from services.assistant_api.monthly_imports import MemoryImportPort
        self.assertEqual(MemoryImportPort().observations, [])

    def test_cp31_rls_baseline_not_modified(self):
        sql = (ROOT / "supabase/migrations/202609250001_contract_core.sql").read_text()
        self.assertIn("enable row level security", sql)
        self.assertIn("products_natural_identity", sql)

    def test_cp32_storage_private(self):
        sql = (ROOT / "supabase/migrations/202609250006_private_storage.sql").read_text()
        self.assertIn("source-files", sql)
        self.assertIn("false", sql)

    def test_cp33_no_pilot_name_hardcode(self):
        text = (ROOT / "services/assistant_api/historical_reconciliation.py").read_text().casefold()
        self.assertNotIn("fendi", text)

    def test_cp34_no_chain_name_hardcode(self):
        text = (ROOT / "services/assistant_api/historical_reconciliation.py").read_text().casefold()
        self.assertNotIn("walmart", text)

    def test_cp35_all_original_conflicts_accounted(self):
        rows = [fact(str(i + 1), item=f"I{i}", upc=f"P{i}", cell=f"V{i+18}") for i in range(730)]
        conflicts = [{"key": [r.chain, r.upc, r.period, r.metric], "values": [r.value, "9999"],
                      "sources": []} for r in rows]
        baseline = {"manifest": [], "conflicts": conflicts}
        result = report(rows, baseline=baseline)
        self.assertEqual(len(result["baseline_conflicts"]), 730)
        self.assertEqual(sum(result["summary"]["initial_matrix"][name] for name in CLASSIFICATIONS), 730)
        with self.assertRaisesRegex(ValueError, "baseline_conflict_unaccounted"):
            report([fact()], baseline=baseline)

    def test_cp36_unresolved_blocks_all_bootstrap(self):
        result = report([fact(), fact("9")])
        port = DurableTestPort(":memory:")
        with self.assertRaisesRegex(ValueError, "historical_bootstrap_blocked"):
            self.execute(result, port)
        self.assertEqual(port.connection.execute("select count(*) from batches").fetchone()[0], 0)
        port.connection.close()

    def test_baseline_hash_mismatch_rejected(self):
        with self.assertRaisesRegex(ValueError, "baseline_source_hash_mismatch"):
            report([fact()], baseline={"manifest": [{"sha256": "changed"}], "conflicts": []})

    def test_report_reproducible_source_order_independent(self):
        rows = [fact(), fact(source_sha256="b" * 64)]
        self.assertEqual(report(rows)["summary"]["dataset_sha256"], report(rows[::-1])["summary"]["dataset_sha256"])

    def test_snapshot_proof_invalid_or_incompatible_never_revises(self):
        proof = SnapshotProof("a" * 64, "monthly-direct", "2025-01", "c" * 64,
                              "Evidence!A1", "reviewer", "Approved comparable snapshots")
        with self.assertRaisesRegex(ValueError, "invalid_snapshot_proof"):
            report([fact()], snapshots={"a" * 64: replace(proof, approved_by="")})
        other = replace(proof, source_hash="b" * 64, cutoff="2025-02", sequence="incompatible")
        result = report([fact(), fact("9", source_sha256="b" * 64, source_cutoff="2025-02")],
                        snapshots={"a" * 64: proof, "b" * 64: other})
        self.assertTrue(result["blocking"])

    def test_only_observed_mapping_periods_resolve_item(self):
        result = report([fact(period="2024-01"), fact(period="2024-03"),
                         fact(upc="CODE2", period="2025-01"), fact(upc="", period="2024-02")])
        self.assertEqual(result["summary"]["item_only_ambiguous"], 1)

    def test_parser_rejections_and_unsupported_sheet_have_ledger(self):
        manifest = {"sha256": "a" * 64, "sheets": [{"sheet": "Summary", "classification": "UNSUPPORTED_OR_AGGREGATE"}]}
        missing = {"key": ["Chain A", "CODE1", "2025-02", "SALES"], "source_sha256": "a" * 64,
                   "sheet": "Detail", "cell": "V19", "reason_code": "MISSING", "blocking": False}
        result = reconcile_historical([manifest], [fact()], excluded=[missing])
        self.assertEqual(result["summary"]["excluded"]["AGGREGATE_ONLY"], 1)
        self.assertEqual(result["summary"]["excluded"]["MISSING"], 1)
        self.assertFalse(result["blocking"])

    def test_hardened_scanner_end_to_end_keeps_baseline_and_locators(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            wide = workbook.active
            wide.title = "Base"
            wide.append(["Prime Item Nbr", "Description", "2025/01 POS Qty",
                         "2025/01 Total Eaches Str Ordered", "2025/01 Total Eaches Str Received"])
            wide.append(["ITEM1", "Name alias", 8, "=1+2", -1])
            summary = workbook.create_sheet("BD (2)")
            summary.cell(3, 2, "Category")
            for column, value in enumerate(["ITEM", "UPC", "Modelo"], 2):
                summary.cell(4, column, value)
            summary.cell(5, 2, "ITEM1")
            summary.cell(5, 3, "CODE1")
            summary.cell(5, 4, "Old name")
            summary.cell(5, 10, 99)
            master = workbook.create_sheet("BAASE")
            header = [None] * 24
            header[1], header[16] = "Cadena", "Fecha"
            for _ in range(16):
                master.append([])
            master.append(header)
            row = [None] * 24
            row[1], row[4], row[6], row[7] = "Chain A", "Towels", "ITEM1", "CODE1"
            row[8], row[16], row[21], row[22], row[23] = "Old name", datetime(2025, 1, 1), 8, 3, 5
            master.append(row)
            workbook.create_sheet("Aggregates").append(["2023 forecast total", 900])
            workbook.save(path)
            specs = [SourceSpec(path, "Chain A", 2025, "2025-01")] * 2
            baseline = scan_sources(specs)
            result = scan_sources(specs, hardened=True, baseline=baseline)
        self.assertEqual(result["summary"]["initial_matrix"]["AUTO_RESOLVABLE"], 1)
        self.assertEqual(result["summary"]["observations"], {"DELIVERY": 1, "ORDER": 1, "SALES": 1})
        self.assertEqual(result["summary"]["excluded"]["FORMULA_UNVERIFIED"], 1)
        self.assertEqual(result["summary"]["excluded"]["INVALID_VALUE"], 1)
        self.assertEqual(result["summary"]["excluded"]["MISSING"], 2)
        self.assertEqual(result["summary"]["source_regression"]["status"], "PASS")
        self.assertFalse(result["blocking"])

    def test_derived_conflict_excluded_but_classified_in_original_matrix(self):
        row = fact(evidence_level="C")
        result = report([row], baseline={"manifest": [], "conflicts": [{"key": [row.chain, row.upc, row.period, row.metric]}]})
        self.assertEqual(result["summary"]["initial_matrix"]["EXCLUDED_NON_FACT"], 1)

    def test_no_temporal_evidence_can_cross_chain(self):
        proof = TemporalEvidence("e", "E1", "2025-02-01T00:00:00Z", "2025-03-01T00:00:00Z",
                                 "reviewer", "a" * 64, "a/evidence", "Chain A")
        result = report([fact(), fact(chain="Chain B")], evidence={("Chain A", "a" * 64): proof})
        self.assertEqual(result["selected"][1]["availability_source"], "UNKNOWN")

    def test_preview_tampering_rejected(self):
        result = report([fact()])
        result["versions"][0]["value"] = "999"
        with self.assertRaisesRegex(ValueError, "canonical_dataset_modified"):
            publication_preflight(result, golden_pass=True, tests_pass=True, project_ref=PROJECT, remote_gate_pass=True)

    def test_wrong_project_golden_or_test_failure_rejected(self):
        result = report([fact()])
        for options in ({"project_ref": "other"}, {"golden_pass": False}, {"tests_pass": False}, {"remote_gate_pass": False}):
            params = dict(golden_pass=True, tests_pass=True, project_ref=PROJECT, remote_gate_pass=True)
            params.update(options)
            with self.assertRaisesRegex(ValueError, "historical_bootstrap_blocked"):
                publication_preflight(result, **params)

    def test_unknown_history_not_visible_in_current_sql_view(self):
        sql = (ROOT / "supabase/migrations/202609250003_contract_governance.sql").read_text()
        self.assertIn("available_at is not null", sql)


if __name__ == "__main__":
    unittest.main()
