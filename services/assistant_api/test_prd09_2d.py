"""PRD 09.2D deterministic tests; private XLSX corpus is scanned separately."""

from __future__ import annotations

import csv
import io
import tempfile
import unittest
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from services.assistant_api.historical_corpus import Candidate, SourceSpec, number, reconcile, scan_sources
from services.assistant_api.monthly_imports import (
    ImportProfileVersion, MemoryImportPort, MonthlyImportWorkflow, parse_monthly,
)


def candidate(value: str, *, period: str = "2025-01", code: str = "7500000000001",
              description: str = "Original", digest: str = "a" * 64,
              cell: str = "V18", chain: str = "Chain A") -> Candidate:
    return Candidate(chain, "ITEM1", code, description, "Towels", period, "SALES",
                     value, digest, "BAASE", cell, "2026-07")


def csv_bytes(rows: list[dict[str, str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["Code", "Variant", "Description", "Category",
                                                 "Period", "Sales", "Order", "Delivery"])
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def row(**overrides: str) -> dict[str, str]:
    result = {"Code": "10001", "Variant": "", "Description": "Towel Blue",
              "Category": "Towels", "Period": "2026-08", "Sales": "10",
              "Order": "12", "Delivery": "9"}
    result.update(overrides)
    return result


def profile(chain: str = "chain-a") -> ImportProfileVersion:
    return ImportProfileVersion("profile-1", chain, 1, {
        "product_code": "Code", "variant_code": "Variant",
        "description": "Description", "category": "Category", "period": "Period",
        "sales": "Sales", "order": "Order", "delivery": "Delivery",
    })


class HistoricalReconciliationTests(unittest.TestCase):
    def test_numeric_missing_error_and_explicit_zero(self) -> None:
        self.assertEqual(number(None), ("MISSING", None))
        self.assertEqual(number("#REF!"), ("INVALID", None))
        self.assertEqual(number("0"), ("VALUE", "0"))
        self.assertEqual(number("1,234"), ("VALUE", "1234"))
        self.assertEqual(number("1,2"), ("INVALID", None))
        self.assertEqual(number(-1), ("INVALID", None))

    def test_duplicate_sources_never_sum(self) -> None:
        report = reconcile([], [candidate("8"), candidate("8", digest="b" * 64)], {}, Counter())
        self.assertEqual(report["summary"]["observations"]["SALES"], 1)
        self.assertEqual(report["summary"]["duplicate_equivalent_sources"], 1)
        self.assertEqual(report["selected"][0]["value"], "8")

    def test_conflicting_snapshots_have_no_winner(self) -> None:
        report = reconcile([], [candidate("8"), candidate("9", digest="b" * 64)], {}, Counter())
        self.assertEqual(report["summary"]["conflicts"], 1)
        self.assertEqual(report["selected"], [])
        self.assertTrue(report["blocking"])

    def test_identity_does_not_use_description(self) -> None:
        report = reconcile([], [candidate("8"), candidate("8", description="New name")], {}, Counter())
        self.assertEqual(report["summary"]["products_unique"], 1)
        self.assertEqual(report["summary"]["description_changes"], 1)

    def test_explicit_item_upc_alias(self) -> None:
        a = candidate("8", code="")
        b = candidate("8")
        report = reconcile([], [a, b], {("Chain A", "ITEM1"): "7500000000001"}, Counter())
        self.assertEqual(report["summary"]["products_unique"], 1)

    def test_prelaunch_zero_is_not_active(self) -> None:
        report = reconcile([], [candidate("0", period="2024-01"),
                                candidate("5", period="2025-01")], {}, Counter())
        self.assertEqual(report["summary"]["not_active"], 1)
        self.assertEqual(report["summary"]["confirmed_zero"], 0)

    def test_zero_after_launch_is_confirmed(self) -> None:
        report = reconcile([], [candidate("1", period="2024-01"),
                                candidate("0", period="2025-01")], {}, Counter())
        self.assertEqual(report["summary"]["confirmed_zero"], 1)
        self.assertEqual(report["selected"][0]["availability_source"], "UNKNOWN")
        self.assertIsNone(report["selected"][0]["available_at"])

    def test_all_zero_has_no_activation_evidence(self) -> None:
        report = reconcile([], [candidate("0")], {}, Counter())
        self.assertEqual(report["summary"]["zero_without_activation_evidence"], 1)
        self.assertTrue(report["blocking"])

    def test_cross_chain_isolation(self) -> None:
        report = reconcile([], [candidate("8"), candidate("10", chain="Chain B")], {}, Counter())
        self.assertEqual(report["summary"]["products_unique"], 2)
        self.assertEqual(report["summary"]["conflicts"], 0)

    def test_realistic_master_xlsx_and_duplicate_file(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "master.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "BAASE"
            for _ in range(16):
                sheet.append([])
            header = [None] * 24
            header[1], header[16] = "Cadena", "Fecha"
            sheet.append(header)
            source = [None] * 24
            source[1], source[4], source[6], source[7] = "Chain A", "Towels", "ITEM1", "7500000000001"
            source[8], source[16] = "Towel", datetime(2025, 1, 1)
            source[21:24] = [5, 6, 7]
            sheet.append(source)
            workbook.save(path)
            report = scan_sources([SourceSpec(path), SourceSpec(path)])
        self.assertEqual(report["summary"]["observations"], {"DELIVERY": 1, "ORDER": 1, "SALES": 1})
        self.assertEqual(report["quality_issues"]["exact_file_duplicate"], 1)
        self.assertEqual(report["manifest"][1]["parsed_observations"], 0)

    def test_wide_and_summary_profiles_join_by_explicit_upc(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wide.xlsx"
            workbook = Workbook()
            wide = workbook.active
            wide.title = "Base"
            wide.append(["Prime Item Nbr", "Prime Item Desc", "2024/01 POS Qty",
                         "2024/01 Total Eaches Str Ordered", "2024/01 Total Eaches Str Received",
                         "2024/02 POS Qty"])
            wide.append(["ITEM1", "Alias description", 4, 6, 5, "=1+1"])
            summary = workbook.create_sheet("BD (2)")
            summary.cell(3, 2, "Cocina")
            for column, value in enumerate(["ITEM", "UPC", "Modelo"], 2):
                summary.cell(4, column, value)
            summary.cell(5, 2, "ITEM1")
            summary.cell(5, 3, "7500000000001")
            summary.cell(5, 4, "Original description")
            summary.cell(5, 10, 4)
            summary.cell(5, 24, 6)
            summary.cell(5, 37, 5)
            workbook.save(path)
            report = scan_sources([SourceSpec(path, "Chain A", 2024, "2024-02")])
        self.assertEqual(report["summary"]["products_unique"], 1)
        self.assertEqual(report["summary"]["description_changes"], 1)
        self.assertEqual(report["summary"]["observations"], {"DELIVERY": 1, "ORDER": 1, "SALES": 1})
        self.assertEqual(report["summary"]["duplicate_equivalent_sources"], 3)
        self.assertEqual(report["quality_issues"]["formula_without_direct_source"], 1)

    def test_unrecognized_workbook_is_not_fabricated_history(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "aggregate.xlsx"
            workbook = Workbook()
            workbook.active.append(["2023 aggregate forecast", 123])
            workbook.save(path)
            report = scan_sources([SourceSpec(path)])
        self.assertTrue(report["blocking"])
        self.assertEqual(report["manifest"][0]["parsed_observations"], 0)
        self.assertEqual(report["quality_issues"]["no_product_actual_profile"], 1)

    def test_master_conflicting_duplicate_rows_are_blocking(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "master.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "BAASE"
            for _ in range(16):
                sheet.append([])
            header = [None] * 24
            header[1], header[16] = "Cadena", "Fecha"
            sheet.append(header)
            for amount in (5, 6):
                source = [None] * 24
                source[1], source[4], source[6], source[7] = "Chain A", "Towels", "ITEM1", "7500000000001"
                source[8], source[16] = "Towel", datetime(2025, 1, 1)
                source[21:24] = [amount, 7, 8]
                sheet.append(source)
            workbook.save(path)
            report = scan_sources([SourceSpec(path)])
        self.assertEqual(report["summary"]["conflicts"], 1)
        self.assertEqual(report["conflicts"][0]["values"], ["5", "6"])


class MonthlyWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.port = MemoryImportPort()
        self.workflow = MonthlyImportWorkflow(self.port, clock=lambda: datetime(2026, 9, 25, 14, tzinfo=timezone.utc))

    def uploaded(self, rows: list[dict[str, str]] | None = None, *, chain: str = "chain-a",
                 filename: str = "monthly.csv"):
        return self.workflow.upload(csv_bytes(rows or [row()]), filename=filename,
                                    profile=profile(chain), uploaded_by="user-1")

    def test_csv_sales_order_delivery_and_timestamp(self) -> None:
        batch = self.uploaded()
        preview = self.workflow.validate(batch.id)
        self.assertEqual(preview["valid_rows"], 1)
        self.assertEqual(self.port.observations, [])
        self.workflow.confirm(batch.id, actor_id="user-1")
        self.assertEqual({r["metric_code"] for r in self.port.observations},
                         {"SALES", "ORDER", "DELIVERY"})
        self.assertTrue(all(r["available_at"] == batch.uploaded_at and
                            r["availability_source"] == "SYSTEM_INGESTION"
                            for r in self.port.observations))

    def test_sha_idempotency_and_private_path(self) -> None:
        first = self.uploaded()
        second = self.uploaded()
        self.assertEqual(first.id, second.id)
        self.assertTrue(first.storage_path.startswith("chain-a/"))
        self.assertEqual(len(self.port.files), 1)

    def test_new_product_and_category_are_chain_scoped(self) -> None:
        a = self.uploaded()
        self.workflow.validate(a.id)
        self.workflow.confirm(a.id, actor_id="user-1")
        b = self.uploaded(chain="chain-b")
        self.workflow.validate(b.id)
        self.workflow.confirm(b.id, actor_id="user-1")
        self.assertEqual(len(self.port.product_rows), 2)
        self.assertNotEqual(self.port.product_rows[("chain-a", "10001", None)]["id"],
                            self.port.product_rows[("chain-b", "10001", None)]["id"])
        self.assertEqual(len(self.port.category_rows), 2)

    def test_description_change_preserves_product_id_and_versions(self) -> None:
        first = self.uploaded()
        self.workflow.validate(first.id)
        self.workflow.confirm(first.id, actor_id="user-1")
        product_id = self.port.product_rows[("chain-a", "10001", None)]["id"]
        second = self.uploaded([row(Description="Towel Azul", Sales="11")])
        preview = self.workflow.validate(second.id)
        self.assertEqual(preview["products_updated"], 1)
        self.workflow.confirm(second.id, actor_id="user-1")
        self.assertEqual(self.port.product_rows[("chain-a", "10001", None)]["id"], product_id)
        versions = [r["version_no"] for r in self.port.observations if r["metric_code"] == "SALES"]
        self.assertEqual(versions, [1, 2])

    def test_preview_does_not_write_facts(self) -> None:
        batch = self.uploaded()
        self.workflow.validate(batch.id)
        self.assertEqual(self.port.product_rows, {})
        self.assertEqual(self.port.category_rows, set())
        self.assertEqual(self.port.observations, [])

    def test_rollback_on_critical_failure(self) -> None:
        batch = self.uploaded([row(), row(Code="10002")])
        self.workflow.validate(batch.id)
        self.port.fail_after = 1
        with self.assertRaisesRegex(RuntimeError, "injected_transaction_failure"):
            self.workflow.confirm(batch.id, actor_id="user-1")
        self.assertEqual(self.port.product_rows, {})
        self.assertEqual(self.port.observations, [])
        self.assertEqual(self.port.batch(batch.id).status, "VALIDATED")

    def test_missing_is_not_zero_and_invalid_formula_rejected(self) -> None:
        accepted, rejected, total = parse_monthly(csv_bytes([row(Sales="", Order="0", Delivery="5"),
                                                             row(Code="10002", Sales="#REF!")]),
                                                  "source.csv", profile())
        self.assertEqual(total, 2)
        self.assertEqual(accepted[0].metrics, {"ORDER": "0", "DELIVERY": "5"})
        self.assertEqual(len(rejected), 1)

    def test_xlsx_and_unverified_formula(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Code", "Variant", "Description", "Category", "Period", "Sales", "Order", "Delivery"])
        sheet.append(["10001", "", "Towel", "Towels", "2026-08", 5, 7, 6])
        sheet.append(["10002", "", "Towel", "Towels", "2026-08", "=1+1", 7, 6])
        buffer = io.BytesIO()
        workbook.save(buffer)
        accepted, rejected, total = parse_monthly(buffer.getvalue(), "source.xlsx", profile())
        self.assertEqual((len(accepted), len(rejected), total), (1, 1, 2))
        self.assertIn("unverified_formula:sales", rejected[0]["errors"])

    def test_invalid_profile_and_format(self) -> None:
        with self.assertRaisesRegex(ValueError, "incomplete_profile"):
            ImportProfileVersion("a", "b", 1, {}).validate()
        with self.assertRaisesRegex(ValueError, "unsupported_import_format"):
            parse_monthly(b"abc", "source.txt", profile())

    def test_bad_filename_and_actor_fail_closed(self) -> None:
        with self.assertRaisesRegex(ValueError, "unsafe_filename"):
            self.uploaded(filename="../source.csv")
        batch = self.uploaded()
        self.workflow.validate(batch.id)
        with self.assertRaisesRegex(ValueError, "confirmation_forbidden"):
            self.workflow.confirm(batch.id, actor_id="other-user")

    def test_hash_mismatch_blocks_validation(self) -> None:
        batch = self.uploaded()
        self.port.files[batch.storage_path] = b"modified"
        with self.assertRaisesRegex(ValueError, "source_hash_mismatch"):
            self.workflow.validate(batch.id)

    def test_previously_present_product_is_never_deleted(self) -> None:
        batch = self.uploaded([row(), row(Code="10002")])
        self.workflow.validate(batch.id)
        self.workflow.confirm(batch.id, actor_id="user-1")
        next_batch = self.uploaded([row(Sales="20")])
        self.workflow.validate(next_batch.id)
        self.workflow.confirm(next_batch.id, actor_id="user-1")
        self.assertIn(("chain-a", "10002", None), self.port.product_rows)


if __name__ == "__main__":
    unittest.main()
