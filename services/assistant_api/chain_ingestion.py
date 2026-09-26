"""Opt-in administrative chain-aware historical ingestion; never runtime cutover.

Private profiles describe source structure. Identity/membership evidence is
independent of literal demand evidence. No retailer/product list or name inference.
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from services.assistant_api.historical_corpus import Candidate, SourceSpec, code, number, sha256_file
from services.assistant_api.historical_reconciliation import reconcile_historical, regression_sample
from services.assistant_api.source_adjudication import fingerprint

ROLES = {"CHAIN_DETAIL", "CHAIN_CALC", "RAW_BASE", "SUMMARY", "FORECAST", "AUXILIARY", "ALTERNATE_REPRESENTATION", "UNKNOWN"}
OUTCOMES = {"RESOLVED_WRONG_CHAIN_GRAIN", "RESOLVED_DUPLICATE_REPRESENTATION", "RESOLVED_MEMBERSHIP",
            "RESOLVED_DETAIL_VS_DERIVED", "STILL_SOURCE_CONFLICT", "STILL_INTERNAL_CONFLICT", "AMBIGUOUS_CHAIN_MEMBERSHIP"}


def unit_key(parent: str, commercial_format: str = "") -> str:
    parent, commercial_format = parent.strip(), commercial_format.strip()
    if not parent:
        return ""
    return f"{parent}::{commercial_format}" if commercial_format and commercial_format != parent else parent


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=") or hasattr(value, "text")


def identifier(value: Any, *, upc: bool = False) -> str:
    if is_formula(value):
        return ""
    text = code(value)
    if not re.fullmatch(r"[A-Za-z0-9_-]+", text) or not any(c.isdigit() for c in text) or text.isdigit() and int(text) == 0:
        return ""
    if upc and (not text.isdigit() or not 8 <= len(text) <= 14):
        return ""
    return text


def dimensions(row: dict[str, Any], profile: SourceSheetProfile) -> tuple[str, str, str]:
    """Explicit row dimensions never fall back to a conflicting sheet assignment."""
    if profile.chain_column:
        parent, fmt = row.get(profile.chain_column), row.get(profile.format_column) if profile.format_column else ""
        if is_formula(parent) or is_formula(fmt) or not isinstance(parent, str) or not parent.strip():
            return "", "", ""
        parent, fmt = parent.strip(), str(fmt or "").strip()
        return parent, fmt, unit_key(parent, fmt)
    return profile.parent_chain, profile.commercial_format, profile.canonical_chain_code or profile.membership_chain_code


@dataclass(frozen=True)
class SourceSheetProfile:
    source_hash: str
    sheet_name: str
    sheet_role: str
    canonical_chain_code: str = ""
    canonical_chain_name: str = ""
    parent_chain: str = ""
    commercial_format: str = ""
    item_column: str = ""
    upc_column: str = ""
    description_column: str = ""
    category_column: str = ""
    period_layout: dict[str, Any] | None = None
    metric_layout: tuple[dict[str, Any], ...] = ()
    metric_source_type: str = "NONE"
    identity_role: str = "NONE"
    membership_role: str = "NONE"
    value_role: str = "NONE"
    notes: str = ""
    chain_column: str = ""
    format_column: str = ""
    header_rows: tuple[int, ...] = ()
    evidence: tuple[dict[str, Any], ...] = ()
    review_required: bool = True
    pair_group: str = ""
    membership_chain_code: str = ""

    def validate(self, source_hash: str, sheet_names: set[str]) -> None:
        if (self.source_hash != source_hash or not re.fullmatch(r"[0-9a-f]{64}", source_hash)
                or self.sheet_name not in sheet_names or self.sheet_role not in ROLES):
            raise ValueError("invalid_source_sheet_profile")
        if (self.canonical_chain_code or self.membership_chain_code) and not self.evidence:
            raise ValueError("chain_assignment_evidence_required")
        if self.sheet_role == "RAW_BASE" and self.canonical_chain_code:
            raise ValueError("raw_base_cannot_be_implicit_chain")
        if self.sheet_role in {"FORECAST", "SUMMARY", "AUXILIARY", "UNKNOWN"} and self.value_role != "NONE":
            raise ValueError("non_fact_sheet_role")
        if any(e.get("source_hash") != source_hash or not e.get("locator") or not e.get("basis") for e in self.evidence):
            raise ValueError("profile_evidence_not_bound_to_source")
        if (self.identity_role not in {"NONE", "EXPLICIT_IDENTIFIERS", "NO_ATOMIC_PRODUCT_ID"}
                or self.membership_role not in {"NONE", "CHAIN_MEMBERSHIP", "SEGMENT_MEMBERSHIP"}
                or self.value_role not in {"NONE", "ACTUAL"}):
            raise ValueError("invalid_profile_role")
        if any(field.get("metric") not in {"SALES", "ORDER", "DELIVERY"}
               or field.get("evidence_level", "A") not in {"A", "C"}
               or not re.fullmatch(r"[A-Z]+", field.get("column", "")) for field in self.metric_layout):
            raise ValueError("invalid_metric_layout")
        for column in (self.item_column, self.upc_column, self.description_column, self.category_column, self.chain_column, self.format_column):
            if column:
                column_index_from_string(column)


@dataclass(frozen=True)
class Membership:
    source_hash: str
    sheet: str
    row: int
    chain: str
    parent_chain: str
    commercial_format: str
    item: str
    upc: str
    role: str
    item_locator: str
    upc_locator: str
    evidence_type: str = "CHAIN_MEMBERSHIP_EVIDENCE"


class MembershipIndex:
    def __init__(self, memberships: list[Membership]):
        self.memberships = memberships
        self.by_item, self.by_upc = defaultdict(list), defaultdict(list)
        self.unresolved_item, self.unresolved_upc = set(), set()
        for entry in memberships:
            if entry.chain and entry.role == "CHAIN_MEMBERSHIP":
                if entry.item:
                    self.by_item[(entry.source_hash, entry.item)].append(entry)
                if entry.upc:
                    self.by_upc[(entry.source_hash, entry.upc)].append(entry)
            elif entry.role == "SEGMENT_MEMBERSHIP":
                if entry.item:
                    self.unresolved_item.add((entry.source_hash, entry.item))
                if entry.upc:
                    self.unresolved_upc.add((entry.source_hash, entry.upc))

    def resolve(self, digest: str, item: str, upc: str, *, explicit_chain: str = "") -> tuple[str, str, str]:
        evidence = self.by_item.get((digest, item), []) if item else self.by_upc.get((digest, upc), [])
        if upc:
            evidence = [entry for entry in evidence if not entry.upc or entry.upc == upc]
        if explicit_chain:
            compatible = [entry for entry in evidence if entry.chain == explicit_chain]
            upcs = {entry.upc for entry in compatible if entry.upc}
            if not upc and len(upcs) > 1:
                return "", "", "AMBIGUOUS_CHAIN_MEMBERSHIP"
            return explicit_chain, upc or (next(iter(upcs)) if len(upcs) == 1 else ""), "EXPLICIT_CHAIN"
        if (digest, item) in self.unresolved_item or (digest, upc) in self.unresolved_upc:
            return "", upc, "AMBIGUOUS_CHAIN_MEMBERSHIP"
        chains = {entry.chain for entry in evidence}
        upcs = {entry.upc for entry in evidence if entry.upc}
        if len(chains) == 1 and len(upcs) <= 1:
            return next(iter(chains)), upc or next(iter(upcs), ""), "UNIQUE_MEMBERSHIP"
        return "", upc, "AMBIGUOUS_CHAIN_MEMBERSHIP" if chains else "MISSING_CHAIN_MEMBERSHIP"


def load_profiles(path: Path) -> list[SourceSheetProfile]:
    values = json.loads(path.read_text(encoding="utf-8"))
    profiles = []
    for value in values:
        for field in ("metric_layout", "header_rows", "evidence"):
            value[field] = tuple(value.get(field, []))
        profiles.append(SourceSheetProfile(**value))
    if len({(p.source_hash, p.sheet_name) for p in profiles}) != len(profiles):
        raise ValueError("duplicate_sheet_profile")
    return profiles


def read_sources(specs: list[SourceSpec]):
    """Sparse literal/formula matrices; no recalc, save, or external-link refresh."""
    books, manifests = {}, []
    for spec in specs:
        digest = sha256_file(spec.path)
        manifests.append({"sha256": digest, "source_file": spec.path.name, "duplicate_relationship": digest if digest in books else None})
        if digest in books:
            continue
        workbook = load_workbook(spec.path, read_only=True, data_only=False)
        sheets = {}
        try:
            for sheet in workbook:
                rows, formulas = {}, set()
                for row_no, cells in enumerate(sheet.iter_rows(), 1):
                    populated = {}
                    for cell in cells:
                        if cell.value is not None:
                            column = get_column_letter(cell.column)
                            populated[column] = cell.value
                            if cell.data_type == "f" or is_formula(cell.value):
                                formulas.add(cell.coordinate)
                    if populated:
                        rows[row_no] = populated
                sheets[sheet.title] = {"rows": rows, "formulas": formulas, "range": sheet.calculate_dimension(),
                                       "max_row": sheet.max_row, "max_column": sheet.max_column}
        finally:
            workbook.close()
        books[digest] = sheets
    return books, manifests


def extract_memberships(books, profiles: list[SourceSheetProfile]) -> list[Membership]:
    result = []
    for profile in profiles:
        if profile.membership_role == "NONE" or not (profile.item_column or profile.upc_column):
            continue
        sheet = books[profile.source_hash][profile.sheet_name]
        for row_no, row in sorted(sheet["rows"].items()):
            if row_no <= min(profile.header_rows, default=0) or row_no in profile.header_rows:
                continue
            item, upc = identifier(row.get(profile.item_column)), identifier(row.get(profile.upc_column), upc=True)
            if not (item or upc):
                continue
            parent, fmt, chain = dimensions(row, profile)
            result.append(Membership(profile.source_hash, profile.sheet_name, row_no, chain,
                                     parent, fmt, item, upc, profile.membership_role,
                                     f"{profile.item_column}{row_no}" if item else "", f"{profile.upc_column}{row_no}" if upc else ""))
    return sorted(result, key=lambda m: (m.source_hash, m.sheet, m.row, m.chain, m.item, m.upc))


def _chain_block(row: Candidate, reason: str) -> dict[str, Any]:
    return {"key": [row.chain, row.upc or row.item, row.period, row.metric],
            "sources": [{"source_sha256": row.source_sha256, "sheet": row.sheet, "cell": row.cell,
                         "item": row.item, "upc": row.upc, "value": row.value, "evidence_level": row.evidence_level}],
            "values": [row.value], "classification": "IDENTITY_BLOCKED", "reason_code": reason,
            "canonical_value": None, "canonical_source": None, "versions": [], "blocking": True,
            "resolution_rule": "NO_RAW_FACT_FANOUT", "notes": "No chain inferred from retailer/filename/description"}


def parse_candidates(books, profiles, memberships):
    index = MembershipIndex(memberships)
    candidates, blocked, excluded, assignments = [], [], [], []
    for profile in profiles:
        sheet = books[profile.source_hash][profile.sheet_name]
        if profile.value_role == "NONE":
            continue
        cutoff = (profile.period_layout or {}).get("cutoff", "")
        for row_no, row in sorted(sheet["rows"].items()):
            if row_no <= min(profile.header_rows, default=0) or row_no in profile.header_rows:
                continue
            item, upc = identifier(row.get(profile.item_column)), identifier(row.get(profile.upc_column), upc=True)
            if not (item or upc):
                continue
            parent, fmt, explicit = dimensions(row, profile)
            chain, mapped_upc, membership_status = index.resolve(profile.source_hash, item, upc, explicit_chain=explicit)
            if profile.chain_column and not explicit:
                chain, membership_status = "", "MISSING_CHAIN_MEMBERSHIP"
            description = str(row.get(profile.description_column) or "").strip()
            category = str(row.get(profile.category_column) or "UNCLASSIFIED").strip()
            for field in profile.metric_layout:
                if not field.get("row_from", 0) <= row_no <= field.get("row_to", sheet.get("max_row", row_no)):
                    continue
                period = field.get("period")
                if field.get("date_column"):
                    when = row.get(field["date_column"])
                    period = when.strftime("%Y-%m") if isinstance(when, (date, datetime)) else None
                if not period or not "2023-01" <= period <= "2026-12" or cutoff and period > cutoff:
                    continue
                cell = f"{field['column']}{row_no}"
                key = [chain or "UNASSIGNED:" + profile.source_hash, mapped_upc or item, period, field["metric"]]
                status, value = number(row.get(field["column"]))
                if cell in sheet["formulas"] or status != "VALUE":
                    excluded.append({"key": key, "source_sha256": profile.source_hash, "sheet": profile.sheet_name,
                                     "cell": cell, "reason_code": "FORMULA_UNVERIFIED" if cell in sheet["formulas"] else
                                     "MISSING" if status == "MISSING" else "INVALID_VALUE", "blocking": False})
                    continue
                candidate = Candidate(key[0], item, mapped_upc, description, category, period, field["metric"],
                                      value, profile.source_hash, profile.sheet_name, cell, cutoff or period,
                                      evidence_level=field.get("evidence_level", "A"))
                assignments.append({"source_sha256": profile.source_hash, "sheet": profile.sheet_name, "cell": cell,
                                    "key": key, "assignment": membership_status,
                                    "parent_chain": parent or profile.parent_chain, "commercial_format": fmt})
                if chain:
                    candidates.append(candidate)
                else:
                    blocked.append(_chain_block(candidate, membership_status))
    return candidates, blocked, excluded, assignments


def reconcile_chain_aware(books, manifests, profiles):
    expected = {(digest, name) for digest, sheets in books.items() for name in sheets}
    supplied = {(p.source_hash, p.sheet_name) for p in profiles}
    if expected != supplied or len(profiles) != len(supplied):
        raise ValueError("every_sheet_requires_profile")
    for profile in profiles:
        profile.validate(profile.source_hash, set(books[profile.source_hash]))
    memberships = extract_memberships(books, profiles)
    candidates, chain_blocked, excluded, assignments = parse_candidates(books, profiles, memberships)
    report = reconcile_historical(manifests, candidates, excluded=excluded)
    # One blocking decision per raw key, retaining every original locator.
    grouped = {}
    for decision in chain_blocked:
        key = tuple(decision["key"])
        if key not in grouped:
            grouped[key] = {**decision, "sources": [], "values": []}
        grouped[key]["sources"].extend(decision["sources"])
        grouped[key]["values"].extend(decision["values"])
        if decision["reason_code"] == "AMBIGUOUS_CHAIN_MEMBERSHIP":
            grouped[key]["reason_code"] = decision["reason_code"]
    for decision in grouped.values():
        decision["sources"] = sorted(decision["sources"], key=lambda s: (s["source_sha256"], s["sheet"], s["cell"]))
        decision["values"] = sorted(set(decision["values"]))
    report["resolution_ledger"].extend(grouped.values())
    report["resolution_ledger"] = sorted(report["resolution_ledger"], key=lambda d: (tuple(d["key"]), fingerprint(d)))
    blockers = [row for row in report["resolution_ledger"] if row["blocking"]]
    report["blocking"] = bool(blockers or not report["selected"])
    report["bootstrap_status"] = "READ_ONLY_CHAIN_MAPPING_REVIEW"
    report["summary"]["blocking_keys"] = len({tuple(row["key"]) for row in blockers})
    report["summary"]["blocking_decisions"] = len(blockers)
    report["summary"]["reason_counts"] = dict(sorted(Counter(row["reason_code"] for row in report["resolution_ledger"]).items()))
    report["summary"]["source_regression"] = regression_sample(report, candidates)
    report["chain_membership_ledger"] = [asdict(m) for m in memberships]
    report["source_assignments"] = sorted(assignments, key=lambda a: (a["source_sha256"], a["sheet"], a["cell"]))
    return report


def trace_old_conflicts(old, new, old_clusters):
    cluster_map = {tuple(key): cluster["cluster_id"] for cluster in old_clusters for key in cluster["keys"]}
    locator = defaultdict(list)
    for decision in new["resolution_ledger"]:
        for source in decision["sources"]:
            if source.get("cell"):
                locator[(source.get("source_sha256"), source.get("sheet"), source["cell"])].append(decision)
    result = []
    for previous in old["resolution_ledger"]:
        if not previous["blocking"]:
            continue
        matched = [decision for source in previous["sources"] for decision in
                   locator[(source["source_sha256"], source["sheet"], source["cell"])]]
        relevant = [d for d in matched if d["key"][2:] == previous["key"][2:] and (d["versions"] or d["blocking"])]
        missing = any(not any(d["key"][2:] == previous["key"][2:] and (d["versions"] or d["blocking"])
                             for d in locator[(s["source_sha256"], s["sheet"], s["cell"])]) for s in previous["sources"] if s["evidence_level"] == "A")
        reasons = {d["reason_code"] for d in relevant}
        still = missing or any(d["blocking"] for d in relevant)
        if missing or reasons & {"AMBIGUOUS_CHAIN_MEMBERSHIP", "MISSING_CHAIN_MEMBERSHIP", "ITEM_ONLY_AMBIGUOUS"}:
            outcome = "AMBIGUOUS_CHAIN_MEMBERSHIP"
        elif still:
            outcome = "STILL_INTERNAL_CONFLICT" if "INTERNAL_SOURCE_DUPLICATE" in reasons else "STILL_SOURCE_CONFLICT"
        elif len({d["key"][0] for d in relevant}) > 1:
            outcome = "RESOLVED_WRONG_CHAIN_GRAIN"
        elif len({d["key"][1] for d in relevant}) > 1:
            outcome = "RESOLVED_MEMBERSHIP"
        elif "DIRECT_VS_SUMMARY" in reasons:
            outcome = "RESOLVED_DETAIL_VS_DERIVED"
        else:
            outcome = "RESOLVED_DUPLICATE_REPRESENTATION"
        result.append({"old_decision_id": "OLD-" + fingerprint(previous["key"])[:12],
                       "old_cluster_id": cluster_map[tuple(previous["key"])],
                       **dict(zip(("old_chain", "old_product", "old_period", "old_metric"), previous["key"])),
                       "new_chain": sorted({d["key"][0] for d in relevant}),
                       "new_product": sorted({d["key"][1] for d in relevant}),
                       "new_classification": sorted({d["classification"] for d in relevant}),
                       "new_resolution": outcome, "still_blocking": still,
                       "source_locators": [[s["source_sha256"], s["sheet"], s["cell"]] for s in previous["sources"]]})
    if len(result) != old["summary"]["blocking_keys"] or len({row["old_decision_id"] for row in result}) != len(result):
        raise ValueError("old_conflict_accounting_mismatch")
    return sorted(result, key=lambda row: row["old_decision_id"])
